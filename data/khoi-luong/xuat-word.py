# Xuất tài liệu Markdown ra Word (A4 ngang). Chạy từ thư mục gốc repo:
#   python3 data/khoi-luong/xuat-word.py docs/phan-tich-khoi-luong.md docs/tai-ve/Phan-tich-khoi-luong-MD-APIC.docx x
# Chuyển Markdown → Word: pandoc rồi chỉnh định dạng (A4 ngang, font, kẻ bảng, tô dòng tiêu đề).
import subprocess, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out, title = sys.argv[1], sys.argv[2], sys.argv[3]
tmp = out + '.tmp.docx'; tmpmd = out + '.tmp.md'
# bỏ các dòng hướng dẫn chạy script (chỉ có ý nghĩa trong repo)
lines = [l for l in open(src, encoding='utf-8') if not l.startswith(('Tính lại:', 'Tạo lại JSON'))]
open(tmpmd, 'w', encoding='utf-8').writelines(lines)
subprocess.run(['pandoc', tmpmd, '-f', 'gfm', '-o', tmp], check=True)
d = Document(tmp)

for sec in d.sections:
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(1.8))
    for m in ('top_margin', 'bottom_margin'): setattr(sec, m, Cm(1.6))

def set_fonts(style, size=None, color=None):
    style.font.name = 'Arial'
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(a), 'Arial')
    rf.set(qn('w:eastAsia'), 'Yu Gothic')
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if size: style.font.size = Pt(size)
    if color: style.font.color.rgb = RGBColor.from_string(color)

for st in d.styles:
    if st.type == 1:  # paragraph
        set_fonts(st)
for name, size in (('Normal', 10), ('Body Text', 10), ('First Paragraph', 10), ('Compact', 9.5),
                   ('Title', 18), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)):
    try: set_fonts(d.styles[name], size, '1F3864' if name.startswith(('Heading', 'Title')) else None)
    except KeyError: pass

def border(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '8EA9C1'); b.append(x)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tblPr.append(b)
    w = tblPr.find(qn('w:tblW'))
    if w is None: w = OxmlElement('w:tblW'); tblPr.append(w)
    w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), '14798')

TOTAL = 14798  # bề rộng vùng in (dxa) khi A4 ngang, lề 1,8 cm
def widths(tbl):
    ncol = len(tbl.columns)
    wt = [0] * ncol
    for row in tbl.rows:
        for j, c in enumerate(row.cells[:ncol]):
            wt[j] = max(wt[j], min(len(c.text), 70))
    wt = [max(w, 10) + 6 for w in wt]
    tot = sum(wt); ws = [int(TOTAL * w / tot) for w in wt]
    grid = tbl._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(ws[j]))
    for row in tbl.rows:
        for j, c in enumerate(row.cells[:ncol]):
            tcPr = c._tc.get_or_add_tcPr(); tw = tcPr.find(qn('w:tcW'))
            if tw is None: tw = OxmlElement('w:tcW'); tcPr.insert(0, tw)
            tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(ws[j]))
    lay = tbl._tbl.tblPr.find(qn('w:tblLayout'))
    if lay is None: lay = OxmlElement('w:tblLayout'); tbl._tbl.tblPr.append(lay)
    lay.set(qn('w:type'), 'fixed')

for tbl in d.tables:
    border(tbl)
    widths(tbl)
    for i, row in enumerate(tbl.rows):
        for cell in row.cells:
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'DCE6F1'); tcPr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0: r.font.bold = True
    # lặp dòng tiêu đề khi bảng sang trang
    trPr = tbl.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)

d.save(out)
import os; os.remove(tmp); os.remove(tmpmd)
print('ok', out)
