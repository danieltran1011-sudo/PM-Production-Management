# Xuất các bảng số liệu phân tích khối lượng nhóm MD ra Excel.
# Chạy từ thư mục gốc repo: python3 data/khoi-luong/xuat-excel.py
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

OUT = 'docs/tai-ve/Phan-tich-khoi-luong-MD-APIC.xlsx'
F = 'Arial'
HEAD = PatternFill('solid', start_color='DCE6F1'); INP = PatternFill('solid', start_color='FFF2CC')
thin = Side(style='thin', color='8EA9C1'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TEN = {'Q.Huy': 'Huy', 'B.Trang': 'Bich Trang'}

wb = Workbook()
def sheet(title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title; return ws
def write(ws, r, c, v, bold=False, fill=None, fmt=None, blue=False, wrap=True):
    x = ws.cell(r, c, v); x.font = Font(name=F, size=10, bold=bold, color='0000FF' if blue else None)
    x.border = BOX; x.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill: x.fill = fill
    if fmt: x.number_format = fmt
    return x
def header(ws, r, cols):
    for j, h in enumerate(cols, 1): write(ws, r, j, h, bold=True, fill=HEAD)
def title(ws, text, note=None):
    ws.cell(1, 1, text).font = Font(name=F, size=13, bold=True, color='1F3864')
    if note: ws.cell(2, 1, note).font = Font(name=F, size=9, italic=True, color='595959')
def widths(ws, ws_):
    for j, w in enumerate(ws_, 1): ws.column_dimensions[L(j)].width = w

# 1. Phân công dự kiến
ws = sheet('Phân công dự kiến', True)
title(ws, 'Phân công dự kiến nhóm MD APIC (quyết định 10/2026)', 'GĐ1: từ nay đến hết mẫu triển lãm 27FW. GĐ2: từ mẫu 1st 28SS (bulk 27SS chạy chuyền).')
rows = [
 ('Cuong', 'HELLY HANSEN (chủ lực); TNF 東京企画 (đến hết GĐ1, bàn giao dần cho Khanh); điều phối xuất hàng chung cả nhóm; quản lý chung may mẫu', 'HELLY HANSEN; điều phối xuất hàng chung; quản lý chung may mẫu', 'NEUTRALWORKS đã chuyển cho Khanh'),
 ('Luan', 'TNF 富山企画 (đầu mối duy nhất với TIJ); một phần việc quản lý (~1 giờ/ngày)', 'Như GĐ1', 'Khanh hỗ trợ ~1/4 việc ở GĐ1, ~1/3 ở GĐ2'),
 ('Khanh', 'NEUTRALWORKS; hỗ trợ TNF 富山企画: báo cáo mã ngoài quy tắc 2 tuần, tiến độ NPL, BL/P/L (cc Luan), dịch một phần tài liệu; theo cùng Cuong các lô TNF 東京企画', 'NEUTRALWORKS; TNF 東京企画 (PPS~đại trà); 1/3 mẫu TNF cùng B.Trang; hỗ trợ TNF 富山企画 (~1/3 việc Luan)', 'Làm đến 16h, có thể tăng ca đến 17h'),
 ('Bich Trang', 'TNF (ALL) mẫu (làm tiếp mẫu triển lãm 27FW); Premium Snow 別注, FOXFIRE, TATRAS, Kasco, YONEX', 'Như GĐ1, chia 1/3 mẫu TNF cho Khanh', ''),
 ('Huy', 'Goldwin (xuyên suốt mẫu và đại trà), TNF 別注 HOSOYA, 141 (hỗ trợ), CANTERBURY, NANGA', 'Giữ nguyên', 'Theo dõi: ~1,39 ở GĐ2'),
 ('Hien', 'Teton Bros, ALLTERRAIN, Mont-Bell, DESCENTE GOLF', 'Giữ nguyên', 'Không giao thêm'),
 ('Nhu', 'TNF 別注 SHIGEMOTO/NAKAO, UTMF, TOKYO MARATHON; việc văn phòng, báo cáo tháng, hóa đơn', 'Giữ nguyên', 'Theo dõi thêm'),
]
header(ws, 4, ['Người', 'GĐ1 – từ nay đến hết mẫu triển lãm 27FW', 'GĐ2 – từ mẫu 1st 28SS', 'Ghi chú'])
for i, r in enumerate(rows, 5):
    for j, v in enumerate(r, 1): write(ws, i, j, v, bold=(j == 1))
widths(ws, [14, 70, 55, 34]); ws.freeze_panes = 'A5'

# 2. Chỉ số tải mùa 27SS (có công thức)
d = json.load(open('data/khoi-luong/hien-tai.json'))
ws = sheet('Chỉ số tải 27SS')
title(ws, 'Chỉ số tải nhóm MD – kỳ hiện tại (bulk 26FW + 27SS, mẫu 27SS)',
      'Ô vàng chữ xanh = số đầu vào (đã nhân hệ số độ khó ALLTERRAIN/DESCENTE × 1,3) và trọng số; có thể sửa. Các cột chỉ số là công thức. 1,0 = trung bình 6 MD làm đủ kỳ.')
write(ws, 4, 1, 'Trọng số', bold=True, fill=HEAD)
for j, (lab, w) in enumerate((('Số mã', .25), ('Ngày chuyền bulk', .25), ('Ngày may mẫu', .25), ('Giờ dịch', .25)), 2):
    write(ws, 3, j, lab, bold=True, fill=HEAD); write(ws, 4, j, w, fill=INP, fmt='0%', blue=True)
write(ws, 4, 6, '=SUM(B4:E4)', fmt='0%'); write(ws, 3, 6, 'Tổng (phải = 100%)', bold=True, fill=HEAD)
cols = ['Người', 'Nhóm', 'Bulk 26FW mã', 'Bulk 26FW SL', 'Bulk 26FW ngày chuyền', 'Bulk 27SS mã', 'Bulk 27SS SL', 'Bulk 27SS ngày chuyền',
        'Mẫu 27SS mã', 'Mẫu 27SS SL', 'Mẫu 27SS ngày may', 'Giờ dịch 27SS (ước)', 'Số brand',
        'Đầu vào: số mã', 'Đầu vào: ngày chuyền bulk', 'Đầu vào: ngày may mẫu', 'Đầu vào: giờ dịch',
        'Chỉ số: số mã', 'Chỉ số: bulk', 'Chỉ số: mẫu', 'Chỉ số: dịch', 'CHỈ SỐ TẢI']
header(ws, 6, cols)
order = ['Hien', 'Cuong', 'Luan', 'Q.Huy', 'B.Trang', 'Nhu', 'Khanh']
r0 = 7; md_rows = []
for i, p in enumerate(order):
    r = r0 + i; v = __import__('collections').defaultdict(float, d[p]); so = 'chi_so' in d[p]
    vals = [TEN.get(p, p), v['nhom'], v['bulk_26FW_ma'], v['bulk_26FW_sl'], v['bulk_26FW_ngay'], v['bulk_27SS_ma'], v['bulk_27SS_sl'], v['bulk_27SS_ngay'],
            v['spl_ma'], v['spl_sl'], v['spl_ngay'], round(v['gio_dich']), v['so_brand'],
            round(v['w_ma'], 1), round(v['w_bulk'], 1), round(v['w_mau'], 1), round(v['w_dich'], 1)]
    for j, x in enumerate(vals, 1):
        write(ws, r, j, x, bold=(j == 1), fill=INP if 14 <= j <= 17 else None, blue=14 <= j <= 17, fmt='#,##0' if isinstance(x, (int, float)) and j < 14 else ('#,##0.0' if j >= 14 else None))
    if so: md_rows.append(r)
last = md_rows[-1]
for r in range(r0, r0 + len(order)):
    for k, (src, dst) in enumerate(zip('NOPQ', 'RSTU')):
        write(ws, r, 18 + k, f'={src}{r}/AVERAGE({src}${md_rows[0]}:{src}${last})', fmt='0.00')
    write(ws, r, 22, f'=SUMPRODUCT($B$4:$E$4,R{r}:U{r})', bold=True, fmt='0.00')
nr = r0 + len(order) + 1
ws.cell(nr, 1, 'Khanh vừa đi làm lại sau nghỉ thai sản (giữa 10/2026): không tính vào mức trung bình. H.Trang thuộc nhóm CAD, không có trong bảng. Nhu ~30% thời gian là việc văn phòng, chưa tính ở đây.').font = Font(name=F, size=9, italic=True)
ws.cell(nr + 1, 1, 'Nguồn: file thực tế BULK 19FW–27SS và mẫu SPL 23FW–27SS (12/10/2026), sheet từng người. Giờ dịch = số mã × giờ dịch/mã của brand.').font = Font(name=F, size=9, italic=True)
widths(ws, [12, 30] + [11] * 15 + [10] * 4 + [11]); ws.freeze_panes = 'C7'
ws.row_dimensions[6].height = 45

# 3. Dự báo 2 giai đoạn (kết quả mô hình)
k = json.load(open('data/khoi-luong/du-bao.json'))
ws = sheet('Dự báo GĐ1-GĐ2')
title(ws, 'Dự báo chỉ số tải theo phương án (11/2026 – 2/2027)',
      'Kết quả từ mô hình data/khoi-luong/du-bao.py (số đã tính, không phải công thức). 1,0 = trung bình nhóm trong giai đoạn. Công suất 9 giờ/ngày, Khanh 8 giờ. Việc giao thêm: Cuong, Luan ~1 giờ/ngày.')
ppl = ['Cuong', 'Luan', 'Khanh', 'B.Trang', 'Q.Huy', 'Hien', 'Nhu']
header(ws, 4, ['Phương án'] + [TEN.get(p, p) for p in ppl])
for i, (name, res) in enumerate(k.items(), 5):
    write(ws, i, 1, name, bold='đã quyết' in name)
    for j, p in enumerate(ppl, 2):
        c = write(ws, i, j, res[p]['chi_so'], fmt='0.00', bold='đã quyết' in name)
        if res[p]['chi_so'] >= 1.4: c.fill = PatternFill('solid', start_color='F8CBAD')
n = 5 + len(k) + 1
for t in ['Tô đỏ: chỉ số ≥ 1,4 (vượt mức tăng ca 1 giờ/ngày, cần chia việc).',
          'GĐ1: mẫu triển lãm 27FW ước bằng SMS 26FW × 40% (chỉ còn theo dõi tiến độ, dịch góp ý). GĐ2: mẫu 1st 28SS ước bằng 1st–3rd 27SS.',
          'Bich Trang ở GĐ2 bị tính thấp: file BULK chưa có bulk 27SS của FOXFIRE, TATRAS, Kasco, YONEX.',
          'Nhu: chỉ tính phần brand; chỉ số thực ≈ chỉ số ÷ 0,7.']:
    ws.cell(n, 1, t).font = Font(name=F, size=9, italic=True); n += 1
widths(ws, [62] + [11] * 7); ws.freeze_panes = 'B5'

# 4. Số mã theo mùa
t = json.load(open('data/khoi-luong/thuc-te.json'))
ws = sheet('Số mã theo mùa')
title(ws, 'Số mã bulk và mẫu theo người, theo mùa', 'Nguồn: sheet từng người trong file BULK và SPL.')
S = ['24FW', '25SS', '25FW', '26SS', '26FW', '27SS']
header(ws, 4, ['Người'] + [f'{s} bulk' for s in S] + [f'{s} mẫu' for s in S])
people = ['Cuong', 'Luan', 'Hien', 'Q.Huy', 'B.Trang', 'Nhu', 'Khanh']
for i, p in enumerate(people, 5):
    write(ws, i, 1, TEN.get(p, p), bold=True)
    for j, s in enumerate(S):
        write(ws, i, 2 + j, sum(r['ma'] for r in t['bulk'] if r['nguoi'] == p and r['mua'] == s), fmt='#,##0')
        write(ws, i, 8 + j, sum(r['ma'] for r in t['spl'] if r['nguoi'] == p and r['mua'] == s), fmt='#,##0')
tr = 5 + len(people)
write(ws, tr, 1, 'Tổng', bold=True)
for j in range(2, 14): write(ws, tr, j, f'=SUM({L(j)}5:{L(j)}{tr-1})', bold=True, fmt='#,##0')
widths(ws, [12] + [10] * 12)

# 5–6. Số liệu gốc
for name, key_, cols_ in (('Thực tế BULK', 'bulk', ['nguoi', 'mua', 'brand', 'sl', 'ma', 'giay', 'ngay', 'ghichu']),
                          ('Thực tế mẫu', 'spl', ['nguoi', 'mua', 'brand', 'loai', 'sl', 'ma', 'giay', 'ngay', 'ghichu'])):
    ws = sheet(name)
    hd = {'nguoi': 'Người', 'mua': 'Mùa', 'brand': 'Brand', 'loai': 'Loại mẫu', 'sl': 'Số lượng', 'ma': 'Số mã', 'giay': 'Giây TB', 'ngay': 'Ngày chuyền / ngày may', 'ghichu': 'Ghi chú'}
    header(ws, 1, [hd[c] for c in cols_])
    for i, r in enumerate(t[key_], 2):
        for j, c in enumerate(cols_, 1):
            v = r.get(c)
            if c == 'giay' and v not in (None, ''): v = round(float(v))
            write(ws, i, j, v, fmt='#,##0' if c in ('sl', 'ma', 'giay', 'ngay') else None, wrap=False)
    ws.auto_filter.ref = f'A1:{L(len(cols_))}{len(t[key_]) + 1}'; ws.freeze_panes = 'A2'
    widths(ws, [10, 8, 32] + [11] * (len(cols_) - 3))

# 7. Phân bổ % thời gian theo bảng nghiệp vụ APIC 2026
a = json.load(open('data/phan-cong/apic-2026-thuong-ky.json'))
ws = sheet('Nghiệp vụ APIC 2026')
nv = [s_['ten'] for s_ in a['nhan_su']]
header(ws, 1, ['STT', 'Phân loại', 'Công đoạn', 'Công việc', 'Nội dung', 'Tần suất'] + nv + ['Tổng'])
for i, x in enumerate(a['cong_viec'], 2):
    for j, v in enumerate([x['stt'], x['phan_loai'], x['cong_doan'], x['cong_viec'], x['noi_dung'], x['tan_suat']], 1): write(ws, i, j, v, wrap=False)
    for j, n_ in enumerate(nv, 7):
        write(ws, i, j, x['phan_bo_cong'].get(n_) or None, fmt='0%')
    write(ws, i, 7 + len(nv), f'=SUM({L(7)}{i}:{L(6 + len(nv))}{i})', fmt='0.00')
er = len(a['cong_viec']) + 2
write(ws, er, 4, 'Tổng mỗi người (= 100%)', bold=True)
for j in range(7, 7 + len(nv)): write(ws, er, j, f'=SUM({L(j)}2:{L(j)}{er-1})', bold=True, fmt='0%')
ws.freeze_panes = 'G2'; widths(ws, [5, 10, 14, 22, 40, 14] + [7] * len(nv) + [7])
ws.cell(er + 2, 1, 'Nguồn: bảng phân tích nghiệp vụ APIC 2026 上期 (cập nhật 1/4/2026), tỷ lệ thời gian khai báo của từng người.').font = Font(name=F, size=9, italic=True)

for ws in wb.worksheets:  # in A4 ngang, vừa bề ngang trang
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT); print('ok', OUT)
