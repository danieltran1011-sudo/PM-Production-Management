# Trích dữ liệu từ Plantable và file mẫu cho mô hình tải theo tháng (ke-hoach-thang.py).
# Chạy từ thư mục gốc repo: python3 data/khoi-luong/trich-ke-hoach.py [PLANTABLE.xlsm] [SAMPLE.xlsx] [từ YYYY-MM-DD] [đến YYYY-MM-DD]
# Kết quả: data/khoi-luong/plantable-10-2026-3-2027.json, data/khoi-luong/mau-2025-2026.json
import sys, json, re, collections, datetime as dt
import openpyxl

PT = sys.argv[1] if len(sys.argv) > 1 else 'data/nguon/APIC-P-Plan-Plantable-2026-20261026.xlsm'
SP = sys.argv[2] if len(sys.argv) > 2 else 'data/nguon/APIC-Sample-20261023.xlsx'
TU = dt.date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else dt.date(2026, 10, 1)
DEN = dt.date.fromisoformat(sys.argv[4]) if len(sys.argv) > 4 else dt.date(2027, 3, 31)
MUA_TOI_THIEU = (26, 1)   # bỏ dòng mùa cũ hơn 26FW (dòng cũ bị đẩy ngày, vd. AP1B-9 dòng 1976–1996)
TEN = {'Như': 'Nhu', 'Luân': 'Luan', 'B Trang': 'B.Trang', 'Trang': 'H.Trang', 'H Trang': 'H.Trang', 'Huy': 'Q.Huy', 'Cường': 'Cuong', 'Khánh': 'Khanh'}
LOAI_PPS = ('PPS', 'RE-PPS', 'BULK', 'MOCK UP', 'MOCKUP', '量産前SPL (TIEN PPS)', '工場控え (LUU)')

def ma(x): return re.sub(r'[\s\.]+$', '', str(x or '').strip().upper())
def mua_so(m):
    m = str(m or '').replace('"', '').strip()
    return (int(m[:2]), 0 if m.endswith('SS') else 1) if m[:2].isdigit() else (0, 0)

# 1) File mẫu: sheet 2025 và 2026 (cột khác nhau)
wb = openpyxl.load_workbook(SP, read_only=True, data_only=True)
mau, pps = [], collections.defaultdict(collections.Counter)
for sn, ci in (('サンプル表2025', dict(p=3, br=5, mua=6, loai=7, ma=8, sl=10, due=11, ok=17)),
               ('サンプル表2026', dict(p=3, br=5, mua=6, loai=7, ma=9, sl=11, due=12, ok=20))):
    if sn not in wb.sheetnames: continue
    for r in wb[sn].iter_rows(min_row=3, max_col=24, values_only=True):
        if not r[ci['p']]: continue
        p = TEN.get(str(r[ci['p']]).strip(), str(r[ci['p']]).strip())
        loai = str(r[ci['loai']] or '').strip().upper()
        if sn == 'サンプル表2026' and r[ci['ma']] and loai in LOAI_PPS: pps[ma(r[ci['ma']])][p] += 1   # chỉ dùng PPS của năm 2026 (đúng người đang phụ trách)
        if not isinstance(r[ci['due']], dt.datetime): continue
        mau.append(dict(nguoi=p, brand=str(r[ci['br']] or '').strip(), mua=str(r[ci['mua']] or '').strip(), loai=loai,
                        sl=r[ci['sl']] if isinstance(r[ci['sl']], (int, float)) else 0, due=r[ci['due']].date().isoformat(),
                        xong=str(r[ci['ok']]).strip() == '●', sheet=sn))

# 2) Gán MD cho mã sản xuất: người theo dõi mẫu PPS của mã; không có thì theo brand
def gan(m, brand, ti, mua):
    b0 = re.sub(r'\s+', '', str(brand or '')).upper()
    if b0 == 'YONEX' and mua_so(mua) >= (27, 0): return 'B.Trang', 'theo brand'   # YONEX từ S27 thuộc B.Trang (bảng 22/7/2026)
    if pps.get(m): c = pps[m]; return max(c, key=c.get), 'file mẫu'
    b = re.sub(r'\s+', '', str(brand or '')).upper()
    if b == 'TNF': p = 'Luan'
    elif b.startswith('TNF'): p = 'Q.Huy' if ti == 'HSY' else 'Nhu'
    elif b.startswith(('HH', 'HELLY')) or b in ('NW', 'NEUTRALWORKS'): p = 'Cuong'
    elif b == 'YONEX': p = 'B.Trang' if mua_so(mua) >= (27, 0) else 'Cuong'
    elif b in ('GOLDWIN', 'NANGA', 'CANTERBURY'): p = 'Q.Huy'
    elif b.startswith(('ALLTERRAIN', 'TETON', 'DESCENT', 'MONT', 'DCC', 'SNOWPEAK')): p = 'Hien'
    elif b in ('FOXFIRE', 'TATRAS', 'KASCO'): p = 'B.Trang'
    elif b in ('UTMF', 'TOKYOMARATHON'): p = 'Nhu'
    elif b == '141': p = 'H.Trang'
    else: p = '?'
    return p, 'theo brand'

# 3) Plantable: sheet lịch theo chuyền
ws = openpyxl.load_workbook(PT, read_only=True, data_only=True)['PLANTABLE (原価) 2026年']
rows, bo = [], []
for i, r in enumerate(ws.iter_rows(min_row=10, max_col=40, values_only=True), 10):
    if not r[8] or not isinstance(r[33], dt.datetime) or not isinstance(r[34], dt.datetime): continue
    s, e = r[33].date(), r[34].date()
    if e < TU or s > DEN: continue
    mua = str(r[0]).replace('"', '').strip()
    x = dict(dong=i, mua=mua, xuong=r[1], line=r[2], tij=str(r[5] or '').strip(), brand=re.sub(r'\s+', ' ', str(r[6] or '')).strip(),
             hd=str(r[7] or '').strip(), ma=ma(r[8]), sl=r[16] if isinstance(r[16], (int, float)) else 0,
             giay=r[28] if isinstance(r[28], (int, float)) else None, start=s.isoformat(), end=e.isoformat())
    if mua_so(mua) < MUA_TOI_THIEU: bo.append(x); continue
    x['nguoi'], x['nguon_gan'] = gan(x['ma'], r[6], x['tij'], mua)
    rows.append(x)
json.dump(dict(rows=rows, loai_bo=dict(ly_do='Dòng mùa cũ hơn 26FW có ngày rơi vào giai đoạn (dòng cũ bị đẩy ngày)', so_dong=len(bo), dong=bo)),
          open('data/khoi-luong/plantable-10-2026-3-2027.json', 'w'), ensure_ascii=False, indent=0)
json.dump(mau, open('data/khoi-luong/mau-2025-2026.json', 'w'), ensure_ascii=False)
print(len(rows), 'lô sản xuất;', len(bo), 'dòng mùa cũ bị loại;', len(mau), 'đơn mẫu;',
      collections.Counter(x['nguon_gan'] for x in rows), collections.Counter(x['nguoi'] for x in rows))
