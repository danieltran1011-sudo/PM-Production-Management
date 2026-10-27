# Tạo lại các file JSON trong data/phan-cong từ file Excel gốc trong data/nguon.
# Chạy từ thư mục gốc repo: python3 data/phan-cong/build.py . data/phan-cong
import json, sys, openpyxl
from openpyxl.utils import get_column_letter as L
src, out = sys.argv[1], sys.argv[2]

def s(v): return None if v is None else str(v).replace('\n',' ').strip() or None
# Nhãn công đoạn bị tách qua 2 dòng ô gộp trong Excel → gộp lại
GOP = {'来客、': '来客・出張対応', '出張対応': '来客・出張対応', '契約管理、': '契約管理・交渉', '契約管理': '契約管理・交渉', '交渉': '契約管理・交渉'}
def cd(v): return GOP.get(v, v)
def num(v): return round(v, 4) if isinstance(v,(int,float)) else v

def analysis(path, sheet, namerow, start, end, factory, year, src_name):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    groups = sorted((c.column, s(c.value)) for c in ws[namerow-1] if s(c.value))
    names = {c.column: s(c.value) for c in ws[namerow] if s(c.value) and s(c.value) != '計'}
    def grp(col): return [g for k,g in groups if k <= col][-1]
    staff = [{"ten": n if n != '-' else 'TIVN本社', "nhom": grp(c)} for c,n in names.items()]
    last = max(names) + 4  # cột 工数 総計
    flag_col = None
    for c in ws[namerow+1]:
        if s(c.value) == '該否': flag_col = c.column
    tasks, cur = [], {}
    for r in range(start, end+1):
        row = {c.column: c.value for c in ws[r]}
        for k in (1,2):
            if s(row.get(k)): cur[k] = s(row[k])
        if not s(row.get(4)) or row.get(3) is None: continue
        alloc = {}
        for c,n in names.items():
            v = row.get(c)
            if isinstance(v,(int,float)) and v: alloc[n if n!='-' else 'TIVN本社'] = round(v,4)
        tasks.append({
            "stt": row.get(3), "phan_loai": cur.get(1), "cong_doan": cd(cur.get(2)) if not str(cur.get(2,'')).startswith('※') else tasks[-1]["cong_doan"] if tasks else None,
            "cong_viec": s(row.get(4)), "noi_dung": s(row.get(5)), "tan_suat": s(row.get(6)),
            "phan_bo_cong": alloc, "tong_cong": round(sum(alloc.values()), 4),
            "tinh_vao_tri_gia_thue": (s(row.get(flag_col)) == '●') if flag_col else None,
        })
    return {"nha_may": factory, "nam_tai_chinh": year, "nguon": src_name, "sheet": sheet.strip(),
            "don_vi_cong": "tỷ lệ thời gian của 1 người (cột mỗi người cộng lại = 1.0)",
            "nhan_su": staff, "cong_viec": tasks}

def old2018(path, src_name):
    ws = openpyxl.load_workbook(path, data_only=True)['APIC BTNB']
    org = sorted((c.column, s(c.value)) for c in ws[2] if s(c.value))
    fac = sorted((c.column, s(c.value)) for c in ws[1] if s(c.value) in ('APIC','BTNB'))
    names = {c.column: s(c.value) for c in ws[4] if c.column >= 14 and s(c.value)}
    staff = []
    for c,n in names.items():
        o = [g for k,g in org if k <= c][-1]
        f = [g for k,g in fac if k <= c]
        staff.append({"ten": n, "to_chuc": o, "nha_may": f[-1] if f and o in ('現地常駐者','TIVN') else None,
                      "loai": s(ws.cell(3,c).value), "cot": L(c)})
    tasks, cur = [], {}
    for r in range(6, 138):
        row = {c.column: c.value for c in ws[r]}
        for k in (1,2):
            if s(row.get(k)): cur[k] = s(row[k])
        if not s(row.get(3)): continue
        ass = {}
        for c,n in names.items():
            v = s(row.get(c))
            if v: ass[f"{n}@{L(c)}"] = v
        tasks.append({"phan_loai": cur.get(1), "cong_doan": cd(cur.get(2)), "cong_viec": s(row.get(3)),
            "noi_dung": s(row.get(4)), "tu": s(row.get(5)), "den": s(row.get(6)),
            "phuong_tien": [m for k,m in ((7,'メール'),(8,'電話'),(9,'作業')) if s(row.get(k))],
            "input": s(row.get(10)), "output": s(row.get(11)),
            "gio_moi_lan": num(row.get(12)), "tan_suat": s(row.get(13)), "phu_trach": ass})
    return {"nguon": src_name, "sheet": "APIC BTNB", "nam": "2018-10",
            "ky_hieu": {"●": "phụ trách chính", "○": "hỗ trợ / tham gia (phía Nhật)", "(●)": "phụ trách khi cần"},
            "nhan_su": staff, "cong_viec": tasks}

N26 = 'data/nguon/Phan-tich-nghiep-vu-TIVN-APIC-2026-thuong-ky.xlsx'
N24 = 'data/nguon/Hop-TIVN-phi-uy-thac-APIC-BTNB-20240130.xlsx'
N18 = 'data/nguon/Phan-cong-nghiep-vu-APIC-BTNB-2018-10.xlsx'
outs = {
 'apic-2026-thuong-ky.json': analysis(N26, '<更新>APIC(2026年度上期) ', 11, 14, 115, 'APIC', '2026 上期 (4–9/2026)', N26),
 'apic-2024.json': analysis(N24, '参考b_APIC(2024年度)', 10, 13, 112, 'APIC', '2024 (5/2024–4/2025)', N24),
 'btnb-2024.json': analysis(N24, '参考b_BTNB(2024年度) ', 11, 14, 113, 'BTNB', '2024 (5/2024–4/2025)', N24),
 'apic-2023.json': analysis(N24, '参考_APIC(2023年度) ', 10, 13, 112, 'APIC', '2023', N24),
 'btnb-2023.json': analysis(N24, '参考_BTNB(2023年度) ', 11, 14, 113, 'BTNB', '2023', N24),
 'phan-cong-2018-10.json': old2018(N18, N18),
}
# Mã hàng cố vấn 天方 phụ trách (26SS)
ws = openpyxl.load_workbook(N26, data_only=True)['<更新>天方氏担当品番_26SS']
rows = []
for r in range(3, ws.max_row+1):
    v = [ws.cell(r,c).value for c in range(1,12)]
    if not v[4]: continue
    rows.append({"dong": s(v[0]), "nguoi_phu_trach_TIVN": s(v[1]), "mua": s(v[2]), "thuong_hieu": s(v[3]), "ma_hang": s(v[4]),
                 "so_hop_dong": s(v[5]), "ke_hoach": v[6], "noi_den": s(v[7]),
                 "etd": v[8].date().isoformat() if hasattr(v[8],'date') else s(v[8]), "da_xuat": v[9], "chenh_lech": v[10]})
outs['ma-hang-co-van-amakata-26ss.json'] = {"nguon": N26, "sheet": "<更新>天方氏担当品番_26SS", "ngay_cap_nhat": "2026-04-01",
    "mo_ta": "Thực tế xuất hàng các mã do cố vấn 天方 (Amakata) phụ trách – TIVN làm rập, nhảy size, tài liệu may", "dong": rows}
for fn, d in outs.items():
    with open(f"{out}/{fn}", 'w', encoding='utf-8') as f: json.dump(d, f, ensure_ascii=False, indent=1)
    print(fn, len(d.get('cong_viec', d.get('dong', []))), 'dòng;', len(d.get('nhan_su', [])), 'người')
