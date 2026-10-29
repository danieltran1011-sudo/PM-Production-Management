# -*- coding: utf-8 -*-
"""APIC 2026年度設備投資 説明資料（費用対効果）ブックを生成する。

元データ: data/nguon/APIC2026年度設備投資予算案_20260403.xlsx（単価・台数・選定理由・議事メモ）
および技術打合せで合意した数値（ロック必要台数、サーボ電力、レンタル実績など）。

実行: python3 dau-tu-thiet-bi-2026/build_workbook.py
出力: dau-tu-thiet-bi-2026/APIC2026年度設備投資_説明資料_費用対効果.xlsx
数式の計算値は LibreOffice で再計算して保存する（Excel では開けば自動計算）。
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "APIC2026年度設備投資_説明資料_費用対効果.xlsx")

FONT = "Meiryo UI"  # 原案ファイルと同じフォント
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F4E78")
F_SEC = Font(name=FONT, size=11, bold=True, color="1F4E78")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_EX = Font(name=FONT, size=10, italic=True, color="808080")
F_RED = Font(name=FONT, size=10, bold=True, color="C00000")

FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_SUB = PatternFill("solid", fgColor="DDEBF7")
FILL_INPUT = PatternFill("solid", fgColor="FFFF99")
FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")
FILL_EX = PatternFill("solid", fgColor="EDEDED")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
CAT_FILL = {
    "必須（安全）": "F8CBAD",
    "必須（顧客要求）": "FCE4D6",
    "必須（生産継続）": "FFF2CC",
    "改善": "E2EFDA",
    "戦略／検証": "DDEBF7",
    "見送り": "EDEDED",
}

THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

USD = '#,##0;(#,##0);"-"'
USD2 = '#,##0.00;(#,##0.00);"-"'
YEN = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
NUM3 = '#,##0.000;(#,##0.000);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()
wb.remove(wb.active)


def sheet(name, widths):
    ws = wb.create_sheet(name)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    return ws


def put(ws, addr, value, font=None, fill=None, fmt=None, align=None, border=True):
    c = ws[addr]
    c.value = value
    c.font = font or F_BASE
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align or WRAP
    if border:
        c.border = BORDER
    return c


def title(ws, text, sub=None):
    put(ws, "A1", text, F_TITLE, border=False, align=Alignment(vertical="center"))
    ws.row_dimensions[1].height = 24
    if sub:
        put(ws, "A2", sub, F_NOTE, border=False, align=Alignment(vertical="top"))


def section(ws, row, text):
    put(ws, f"A{row}", text, F_SEC, border=False, align=Alignment(vertical="center"))
    ws.row_dimensions[row].height = 20


def header(ws, row, labels, start="A"):
    col0 = ord(start) - 64
    for i, lab in enumerate(labels):
        col = chr(64 + col0 + i) if col0 + i <= 26 else "A" + chr(64 + col0 + i - 26)
        put(ws, f"{col}{row}", lab, F_HEAD, FILL_HEAD, align=CENTER)
    ws.row_dimensions[row].height = 32


def text_table(ws, row, labels, rows, heights=None, start="A"):
    """見出し＋文字列の表。戻り値は次の空き行。"""
    header(ws, row, labels, start)
    col0 = ord(start) - 64
    for r, data in enumerate(rows, start=row + 1):
        for i, v in enumerate(data):
            put(ws, f"{chr(64 + col0 + i)}{r}", v)
        if heights:
            ws.row_dimensions[r].height = heights
    return row + len(rows) + 1


def note(ws, addr, text):
    put(ws, addr, text, F_NOTE, border=False, align=Alignment(wrap_text=False, vertical="top"))


def inp(ws, addr, value, fmt=None, fill=True, comment=None):
    c = put(ws, addr, value, F_INPUT, FILL_INPUT if fill else None, fmt)
    if comment:
        c.comment = Comment(comment, "APIC")
    return c


def link(ws, addr, formula, fmt=None):
    return put(ws, addr, formula, F_LINK, fmt=fmt)


def calc(ws, addr, formula, fmt=None, bold=False, fill=None):
    return put(ws, addr, formula, F_BOLD if bold else F_BASE, fill, fmt)


# ---------------------------------------------------------------- シート名
S_COVER = "00_表紙"
S_PRM = "01_前提条件"
S_SUM = "02_投資総括"
S_QA = "03_要確認事項"
S_CMP = "04_機種別比較"
S_LOCK = "05_ロック必要台数"
S_SERVO = "06_サーボ電力"
S_PS = "07_省人効果"
S_OPS = "08_PS運用設計"
S_SPOT = "09_スポットプレス"
S_FLAT = "10_平プレス検証"
S_FELT = "11_縮絨試験"
S_RENT = "12_レンタル費"
S_ROI = "13_費用対効果PL"
S_CHK = "14_原案チェック"


def ref(sh, addr):
    return f"'{sh}'!{addr}"


# 前提条件のセル（他シートから参照）
RATE = ref(S_PRM, "$C$5")
DAYS = ref(S_PRM, "$C$6")
HOURS = ref(S_PRM, "$C$7")
SECS = ref(S_PRM, "$C$9")
ELEC = ref(S_PRM, "$C$10")
YEARS = ref(S_PRM, "$C$11")
APIC = ref(S_PRM, "$C$12")
CAP = ref(S_PRM, "$C$13")
LABOR = ref(S_PRM, "$C$14")
LINES = ref(S_PRM, "$C$15")
PEAK = ref(S_PRM, "$C$16")

# 表紙・目次は最後に作成するので先にシートだけ確保
cover = sheet(S_COVER, {"A": 4, "B": 26, "C": 60, "D": 22, "E": 22})

# ================================================================ 01 前提条件
ws = sheet(S_PRM, {"A": 34, "B": 12, "C": 16, "D": 80})
title(ws, "前提条件（全シート共通）", "青字＋黄色セルが入力値。変更すると全シートが再計算される。")
header(ws, 3, ["項目", "単位", "値", "出所・備考"])
params = [
    (5, "為替レート", "円/USD", 160, "#,##0", "原案P30（円換算に使用）。原案O26は150円 → 提出前にTIJとレートを統一"),
    (6, "年間稼働日", "日", 295, "#,##0", "技術打合せでの合意値"),
    (7, "1日稼働時間", "時間", 8, "#,##0.0", "技術打合せでの合意値"),
    (8, "有効稼働率", "%", 1.0, PCT, "工程秒が標準時間（余裕込み）なら100%。正味時間で測る場合は工場の実効率を入力"),
    (10, "電気単価", "USD/kWh", 0.08, "0.000", "技術打合せでの合意値"),
    (11, "償却年数", "年", 7, "#,##0", "原案U列（=金額÷7）と同じ"),
    (12, "APIC負担額", "USD", 50000, USD, "新規条件：APICが投資額のうち50,000USDを負担（対象機種・資産保有・償却処理は要確認）"),
    (13, "投資額上限目標（Suntool除く）", "円", 40000000, YEN, "佐藤さん希望：Suntoolを除き概ね4,000万円以下。20260403版リストにSuntoolは含まれていない"),
    (14, "1人当たり年間人件費", "USD/人年", None, USD, "APIC人事データで入力。未入力の間は省人効果を『人工』のみで表示し、金額効果は0とする"),
    (15, "通常稼働ライン数", "ライン", 31, "#,##0", "APIC通常31ライン"),
    (16, "ピークライン数", "ライン", 35, "#,##0", "繁忙期34〜35ライン（軽量品6,000秒未満は32人ラインを16人×2に分割）→ 上限35で試算"),
]
for r, item, unit, val, fmt, src in params:
    put(ws, f"A{r}", item)
    put(ws, f"B{r}", unit, align=CENTER)
    inp(ws, f"C{r}", val, fmt)
    put(ws, f"D{r}", src)
put(ws, "A9", "1人年間有効稼働秒")
put(ws, "B9", "秒", align=CENTER)
calc(ws, "C9", "=C6*C7*3600*C8", "#,##0", bold=True)
put(ws, "D9", "年間稼働日×1日稼働時間×3600×有効稼働率（省人工数の分母）")

# ================================================================ 02 投資総括
ws = sheet(S_SUM, {"A": 6, "B": 26, "C": 22, "D": 16, "E": 7, "F": 10, "G": 8, "H": 11,
                   "I": 9, "J": 11, "K": 13, "L": 10, "M": 46, "N": 11, "O": 11, "P": 12, "Q": 16})
title(ws, "投資総括：原案（優先①）と調整案",
      "調整案台数（I列）とAPIC負担（N列）が入力欄。スポットプレスの台数は09シートの計算結果を自動参照。単価は原案L列（VA後）、L列が空欄の機種はK列（概算）。")
header(ws, 4, ["原案No.", "設備名", "メーカー・品番", "投資区分", "優先順位", "単価(USD)", "原案①台数",
               "原案①金額(USD)", "調整案台数", "調整案金額(USD)", "調整案金額(円)", "増減(USD)",
               "調整理由・判断", "APIC負担(USD)", "TI負担(USD)", "TI年間償却(円)", "根拠シート"])
ws.freeze_panes = "C5"

items = [
    # (no, 名称, 品番, 区分, 優先, 単価, 原案①, 調整案(数値 or None=スポット), 理由, 根拠)
    ("1", "2本針オーバーロック（サーボ・自動糸切り）", "Juki MO-6814", "必須（生産継続）", 3, 980, 5, 5,
     "ピーク不足26台（05）。故障4台の更新＋純増1台。残る不足は第2工場借用で対応", S_LOCK),
    ("2", "2本針インターロック（サーボ・自動糸切り）", "Juki MO-6816", "必須（生産継続）", 3, 1150, 2, 2,
     "ピーク不足6台（05）。2台購入、残りは借用で対応", S_LOCK),
    ("3", "眠り穴かがりミシン", "Hikari HBH-1790E-S", "見送り", None, 2376, 0, 0,
     "原案でも優先②のみ。2026年度は計上しない", ""),
    ("4", "パターンシーマ 中型 60×40", "Hikari HMS-6040", "改善", 4, 4450, 3, 3,
     "GW系ポケットファスナーのPS指示。運用設計（08）とセットで導入、効果は07で実測", S_PS),
    ("5", "パターンシーマ 大型 130×95（360度回転）", "Hikari HMS-13095A/RND", "改善", 4, 6500, 6, 4,
     "6→4台に段階導入。数値で確認できる効果はレンタル3,016.54＋キルティング外注約3,000USD/年。省人実測後に追加判断", S_PS),
    ("6", "本縫針送りミシン", "DUMA DM1968L-NF", "必須（生産継続）", 5, 880, 10, 10,
     "各ライン1台必要（31台）に対し現有17台。レンタル削減額は2025実績で再確認（14参照）", S_CHK),
    ("7", "シームシーリング Spec-D（ロングアーム）", "H&H AI-002 spec-D", "改善", 7, 9700, 3, 2,
     "3→2台。2名体制→1名化の省人効果を07で実測し、効果確認後に追加", S_PS),
    ("8", "シームシーリング Spec-IJ（曲線部）", "H&H AI-129 spec-IJ", "戦略／検証", 9, 6800, 1, 1,
     "評価用1台（量産ライン配置、剥離不良率で評価）", S_CMP),
    ("9", "温冷プレス 60×12cm", "H&H CSL-201", "必須（生産継続）", 6, 4150, 6, 6,
     "現有30×10cm 26台で不足。工程秒約半減（原案記載）→ 対象数量で効果算出", S_ROI),
    ("10", "平プレス 125×83cm", "H&H CS-691", "改善", 8, 13000, 2, 1,
     "2→1台。精度検証方法（10）で効果確認後に追加", S_FLAT),
    ("11", "スポットプレス 30×10cm（安全装置付）", "PTTG MEN/PTTG 300X100", "必須（安全）", 1, 764, 10, None,
     "安全装置なし旧型を全数入替。台数は09の計算結果（暫定：必要45台−新型現有15台）", S_SPOT),
    ("12", "Nawon温冷プレス（フロントジッパー）", "Nawon HTM-8749", "必須（顧客要求）", 2, 19980, 2, 2,
     "26FW GW要請（SW1課より4種導入要望）", S_CMP),
    ("13", "Nawon温冷プレス（袖）", "Nawon HTM-8739", "必須（顧客要求）", 2, 19980, 1, 1,
     "同上", S_CMP),
    ("14", "Nawon温冷プレス（ヘムライン）", "Nawon HTM-8709", "必須（顧客要求）", 2, 19980, 1, 1,
     "同上。ストレート型はH&H CSL-207（10,500USD）で代替可能性あり → 追加削減オプション", S_CMP),
    ("15", "Nawon温冷プレス（カーブ）", "Nawon HTM-8601", "必須（顧客要求）", 2, 18333, 1, 1,
     "同上（2025年度1台導入済）", S_CMP),
    ("16", "自動延反機", "NCA ES-8", "見送り", None, 50980, 0, 0, "2/25打合せで26年度見送り", ""),
    ("17", "ダウン充填機", "Shenghong 1184RM10", "見送り", None, 9800, 0, 0,
     "3/25打合せで見送り（現有4台で最大10ライン対応可）", ""),
    ("18", "超音波溶着機 35Hz", "H&H US-512", "戦略／検証", 10, 6300, 1, 1,
     "トライアル（サンプル試作で評価、営業提案の幅を広げる）", S_CMP),
    ("18※", "テーピングマシン（比較検証用）", "Nawon HTM-5588 SSF-LDi", "見送り", None, 11100, 0, 0,
     "原案でも優先②のみ。※原案でNo.18が重複", ""),
    ("19", "テーピングマシン（比較検証用）", "Nawon HTM-3833 SSF-Ldi", "見送り", None, 9130, 0, 0, "原案でも優先②のみ", ""),
    ("20", "刺繍機（20頭）", "TAJIMA TMCP-V0920F", "見送り", None, 81000, 0, 0, "原案でも優先②のみ", ""),
]
SUM_ROW = {}  # 原案No. → 行
first = 5
dv = DataValidation(type="list", formula1='"必須（安全）,必須（顧客要求）,必須（生産継続）,改善,戦略／検証,見送り"',
                    allow_blank=True)
ws.add_data_validation(dv)
for i, (no, name, model, cat, prio, price, q0, q1, reason, src) in enumerate(items):
    r = first + i
    SUM_ROW[no] = r
    put(ws, f"A{r}", no, align=CENTER)
    put(ws, f"B{r}", name)
    put(ws, f"C{r}", model)
    c = put(ws, f"D{r}", cat, fill=PatternFill("solid", fgColor=CAT_FILL[cat]), align=CENTER)
    dv.add(c)
    put(ws, f"E{r}", prio if prio else "-", align=CENTER)
    put(ws, f"F{r}", price, F_INPUT, fmt=USD)
    put(ws, f"G{r}", q0, F_INPUT, fmt="#,##0", align=CENTER)
    calc(ws, f"H{r}", f"=F{r}*G{r}", USD)
    if q1 is None:
        link(ws, f"I{r}", "=" + ref(S_SPOT, "$C$20"), "#,##0")
    else:
        inp(ws, f"I{r}", q1, "#,##0")
    calc(ws, f"J{r}", f"=F{r}*I{r}", USD)
    calc(ws, f"K{r}", f"=J{r}*{RATE}", YEN)
    calc(ws, f"L{r}", f"=J{r}-H{r}", USD)
    put(ws, f"M{r}", reason)
    inp(ws, f"N{r}", None, USD)
    calc(ws, f"O{r}", f"=J{r}-N{r}", USD)
    calc(ws, f"P{r}", f"=O{r}*{RATE}/{YEARS}", YEN)
    put(ws, f"Q{r}", src)
    ws.row_dimensions[r].height = 42
last = first + len(items) - 1
T = last + 1  # 合計行
put(ws, f"A{T}", "合計", F_BOLD, FILL_TOTAL)
ws.merge_cells(f"A{T}:F{T}")
for col in "GHIJKLNOP":
    fmt = YEN if col in "KP" else ("#,##0" if col in "GI" else USD)
    calc(ws, f"{col}{T}", f"=SUM({col}{first}:{col}{last})", fmt, bold=True, fill=FILL_TOTAL)
for col in "MQ":
    put(ws, f"{col}{T}", "", fill=FILL_TOTAL)
SUM_T = T

r = T + 2
section(ws, r, "総額チェック")
header(ws, r + 1, ["", "項目", "USD", "円", "判定・備考"])
rows = [
    ("原案（優先①）合計", f"=H{T}", f"=H{T}*{RATE}", "原案ファイルのN25（267,343USD）と一致"),
    ("調整案 合計", f"=J{T}", f"=K{T}", ""),
    ("投資額上限目標（Suntool除く）", f"={CAP}/{RATE}", f"={CAP}", "前提条件C13"),
    ("上限との差（＋＝余裕）", f"=C{r+4}-C{r+3}", f"=D{r+4}-D{r+3}",
     f'=IF(D{r+5}>=0,"上限内","上限超過：追加削減オプションを検討")'),
    ("APIC負担（前提）", f"={APIC}", f"={APIC}*{RATE}",
     f'=IF(N{T}={APIC},"機種別配分 OK","機種別配分 未確定（N列に入力：未割当 "&TEXT({APIC}-N{T},"#,##0")&" USD）")'),
    ("TI負担（調整案−APIC負担）", f"=C{r+3}-C{r+6}", f"=D{r+3}-D{r+6}", "APIC負担分の資産保有・償却処理は要確認（13参照）"),
]
for k, (lab, u, y, rem) in enumerate(rows):
    rr = r + 2 + k
    put(ws, f"B{rr}", lab, F_BOLD)
    calc(ws, f"C{rr}", u, USD, bold=True, fill=FILL_KEY if k in (1, 3) else None)
    calc(ws, f"D{rr}", y, YEN, bold=True, fill=FILL_KEY if k in (1, 3) else None)
    put(ws, f"E{rr}", rem)
ws.merge_cells(f"E{r+1}:H{r+1}")
for k in range(len(rows)):
    ws.merge_cells(f"E{r+2+k}:H{r+2+k}")
CHK_ROW_TOTAL = r + 3   # 調整案 合計の行
CHK_ROW_GAP = r + 5

r = r + 2 + len(rows) + 1
section(ws, r, "投資区分別（調整案）")
header(ws, r + 1, ["", "投資区分", "調整案金額(USD)", "構成比", "考え方"])
cats = [
    ("必須（安全）", "安全リスクの解消。ROIではなく安全で判断"),
    ("必須（顧客要求）", "顧客指示・要請への対応（受注維持）"),
    ("必須（生産継続）", "機材不足の解消（借用・レンタル依存からの脱却）"),
    ("改善", "省人・工数削減・レンタル/外注削減・品質。数値効果で判断し段階導入"),
    ("戦略／検証", "将来案件・トライアル。少数導入して評価"),
    ("見送り", "2026年度は計上しない"),
]
for k, (cat, idea) in enumerate(cats):
    rr = r + 2 + k
    put(ws, f"B{rr}", cat, fill=PatternFill("solid", fgColor=CAT_FILL[cat]))
    calc(ws, f"C{rr}", f'=SUMIF($D${first}:$D${last},B{rr},$J${first}:$J${last})', USD)
    calc(ws, f"D{rr}", f"=IF($J${T}=0,0,C{rr}/$J${T})", PCT)
    put(ws, f"E{rr}", idea)
    ws.merge_cells(f"E{rr}:H{rr}")
ws.merge_cells(f"E{r+1}:H{r+1}")

r = r + 2 + len(cats) + 1
section(ws, r, "追加削減オプション（上限超過時・TIJ指摘時の選択肢）")
header(ws, r + 1, ["", "内容", "削減額(USD)", "削減額(円)", "留意点"])
opts = [
    ("No.14 HTM-8709 → H&H CSL-207（ストレート型）", f"=F{SUM_ROW['14']}-10500",
     "ストレート型は性能差が小さい（2/25打合せ）。GW要請は『Nawon 4種』のため営業確認が必要"),
    ("No.9 温冷プレス 6台 → 4台", f"=F{SUM_ROW['9']}*2", "不足台数の再計算が前提"),
    ("No.18 US-512 を優先②へ", f"=F{SUM_ROW['18']}", "トライアルを次年度へ延期"),
    ("No.8 Spec-IJ を優先②へ", f"=F{SUM_ROW['8']}", "曲線部シーム剥離の検証が遅れる"),
]
for k, (lab, f, rem) in enumerate(opts):
    rr = r + 2 + k
    put(ws, f"B{rr}", lab)
    calc(ws, f"C{rr}", f, USD)
    calc(ws, f"D{rr}", f"=C{rr}*{RATE}", YEN)
    put(ws, f"E{rr}", rem)
    ws.merge_cells(f"E{rr}:H{rr}")
    ws.row_dimensions[rr].height = 30
ws.merge_cells(f"E{r+1}:H{r+1}")
note(ws, f"B{r+2+len(opts)}", "※単価10,500USD（CSL-207）は原案T17の比較情報。")


def J(no):
    return ref(S_SUM, f"$J${SUM_ROW[no]}")


def I(no):
    return ref(S_SUM, f"$I${SUM_ROW[no]}")


# ================================================================ 05 ロック必要台数
ws = sheet(S_LOCK, {"A": 40, "B": 16, "C": 16, "D": 60})
title(ws, "ロックミシン 必要台数と不足（No.1・No.2）",
      "不足は試算上だけではなく、第2工場からの借用および外部レンタルで実運用上すでに顕在化している。")
section(ws, 3, "1. 品種別 標準使用台数（1ライン当たり）")
text_table(ws, 4, ["品種", "オーバーロック", "インターロック", "主な用途"], [
    ("通常品", 3, 2, "OL：パーツ縫製・ポケット袋布・収納袋／IL：袖下・脇・フード"),
    ("Seam Jacket", 3, 2, "同上"),
    ("Down Jacket", 3, 2, "同上"),
    ("ウィンドブレーカー・1枚もの軽量品", 4, 2, "軽量品は全パーツの端処理をOLで行うため1台多い"),
])
section(ws, 10, "2. 前提")
header(ws, 11, ["項目", "値", "", "備考"])
lock_in = [
    (12, "通常稼働ライン数", "=" + LINES, "#,##0", "前提条件"),
    (13, "ピークライン数", "=" + PEAK, "#,##0", "前提条件"),
    (14, "OL3台ラインの構成比", 0.8, PCT, "製品構成の想定（通常品・Seam・Down）"),
    (15, "OL4台ラインの構成比", "=1-B14", PCT, "軽量品"),
]
for rr, lab, v, fmt, rem in lock_in:
    put(ws, f"A{rr}", lab)
    if isinstance(v, str) and v.startswith("='") or (isinstance(v, str) and LINES in v) or (isinstance(v, str) and PEAK in v):
        link(ws, f"B{rr}", v, fmt)
    elif isinstance(v, str):
        calc(ws, f"B{rr}", v, fmt)
    else:
        inp(ws, f"B{rr}", v, fmt)
    put(ws, f"D{rr}", rem)
put(ws, "A16", "平均OL台数／ライン", F_BOLD)
calc(ws, "B16", "=B14*3+B15*4", "0.0", bold=True)
put(ws, "D16", "80%×3台＋20%×4台＝3.2台（原案R4の『平均4台』は過大 → 修正）")
put(ws, "A17", "IL台数／ライン")
inp(ws, "B17", 2, "#,##0")
put(ws, "D17", "全品種2台")

section(ws, 19, "3. 必要台数・実稼働・不足")
header(ws, 20, ["項目", "オーバーロック", "インターロック", "計算式・備考"])
lock_rows = [
    (21, "現有台数", (102, F_INPUT), (68, F_INPUT), "原案S4・S5"),
    (22, "老朽化・重大故障（修理復旧不可、実稼働対象外）", (4, F_INPUT), (0, F_INPUT),
     "OL4台は修理復旧不可。ILは0〜1台（実稼働63〜64台）→ 要確認"),
    (23, "サンプルライン固定", (12, F_INPUT), (4, F_INPUT), ""),
    (24, "生産用 実稼働可能台数", ("=B21-B22-B23", None), ("=C21-C22-C23", None), "現有−故障−サンプル固定"),
    (25, "通常期 必要台数（31ライン）", ("=ROUNDUP(B12*B16,0)", None), ("=B12*B17", None), "ライン数×平均台数（端数切上げ）"),
    (26, "ピーク 必要台数（35ライン）", ("=ROUNDUP(B13*B16,0)", None), ("=B13*B17", None), ""),
    (27, "通常期 不足", ("=MAX(0,B25-B24)", None), ("=MAX(0,C25-C24)", None), ""),
    (28, "ピーク 不足（実質不足）", ("=MAX(0,B26-B24)", None), ("=MAX(0,C26-C24)", None), "この不足を借用・レンタルで埋めている"),
]
for rr, lab, ol, il, rem in lock_rows:
    put(ws, f"A{rr}", lab, F_BOLD if rr == 28 else F_BASE)
    for col, (v, f) in (("B", ol), ("C", il)):
        if f is F_INPUT:
            inp(ws, f"{col}{rr}", v, "#,##0")
        else:
            calc(ws, f"{col}{rr}", v, "#,##0", bold=rr == 28, fill=FILL_KEY if rr == 28 else None)
    put(ws, f"D{rr}", rem)

section(ws, 30, "4. 実運用での顕在化（証跡）")
header(ws, 31, ["項目", "台数", "", "備考"])
put(ws, "A32", "繁忙期 APIC第2工場から借用")
inp(ws, "B32", 10, "#,##0")
put(ws, "D32", "10〜12台（下限で計上）。先方の受注増で継続借用は困難になる見込み。機種内訳（OL/IL）は要確認")
put(ws, "A33", "外部レンタル実績（2024年10〜12月）")
inp(ws, "B33", 10, "#,##0")
put(ws, "D33", "約10台。金額はD/Nから12_レンタル費へ入力")

section(ws, 35, "5. 購入後の不足")
header(ws, 36, ["項目", "オーバーロック", "インターロック", "備考"])
put(ws, "A37", "調整案 購入台数")
link(ws, "B37", "=" + I("1"), "#,##0")
link(ws, "C37", "=" + I("2"), "#,##0")
put(ws, "D37", "02_投資総括から参照")
put(ws, "A38", "うち 更新分（故障機の置換）")
calc(ws, "B38", "=MIN(B37,B22)", "#,##0")
calc(ws, "C38", "=MIN(C37,C22)", "#,##0")
put(ws, "D38", "")
put(ws, "A39", "うち 純増分")
calc(ws, "B39", "=B37-B38", "#,##0")
calc(ws, "C39", "=C37-C38", "#,##0")
put(ws, "D39", "")
put(ws, "A40", "購入後 ピーク不足", F_BOLD)
calc(ws, "B40", "=MAX(0,B28-B37)", "#,##0", bold=True)
calc(ws, "C40", "=MAX(0,C28-C37)", "#,##0", bold=True)
put(ws, "D40", "残りは第2工場借用で対応（借用不可の場合は原案優先②5台＋3台を条件付き追加枠とする）")
note(ws, "A42", "説明文案：『不足は試算上だけではなく、繁忙期に第2工場から10〜12台を借用し、2024年10〜12月には約10台を外部レンタルしており、"
     "実運用上すでに顕在化している。』")

# ================================================================ 06 サーボ電力
ws = sheet(S_SERVO, {"A": 38, "B": 14, "C": 14, "D": 60})
title(ws, "ロックミシン サーボモーター vs クラッチモーター",
      "電力削減は副次効果。主目的は機材不足の解消と、即停止による縫いズレ低減（品質・操作性）。")
section(ws, 3, "1. 消費電力の前提")
header(ws, 4, ["項目", "クラッチ", "サーボ", "備考"])
servo = [
    (5, "定格出力 (W)", 400, 550, "サーボは定格が大きいがペダル操作時のみ回転"),
    (6, "縫製時 (W)", 350, 200, ""),
    (7, "非縫製時 (W)", 250, 2, "クラッチは縫っていない間も空転、サーボは待機電力のみ"),
]
for rr, lab, a, b, rem in servo:
    put(ws, f"A{rr}", lab)
    inp(ws, f"B{rr}", a, "#,##0")
    inp(ws, f"C{rr}", b, "#,##0")
    put(ws, f"D{rr}", rem)
put(ws, "A8", "1日稼働時間 (h)")
link(ws, "B8", "=" + HOURS, "0.0")
calc(ws, "C8", "=B8", "0.0")
put(ws, "D8", "前提条件")
put(ws, "A9", "うち縫製時間 (h)")
inp(ws, "B9", 4, "0.0")
calc(ws, "C9", "=B9", "0.0")
put(ws, "D9", "8時間中4時間縫製 → クラッチ2.400kWh/日・サーボ0.808kWh/日（従来試算と一致）")
put(ws, "A10", "非縫製時間 (h)")
calc(ws, "B10", "=B8-B9", "0.0")
calc(ws, "C10", "=C8-C9", "0.0")
put(ws, "D10", "")
section(ws, 12, "2. 削減効果")
header(ws, 13, ["項目", "値", "", "計算式"])
srows = [
    (14, "クラッチ 消費電力 (kWh/日)", "=(B6*B9+B7*B10)/1000", NUM3, "(縫製W×縫製h＋非縫製W×非縫製h)÷1000"),
    (15, "サーボ 消費電力 (kWh/日)", "=(C6*C9+C7*C10)/1000", NUM3, ""),
    (16, "削減 (kWh/日・台)", "=B14-B15", NUM3, ""),
    (17, "年間削減 (kWh/年・台)", f"=B16*{DAYS}", NUM2, "×年間稼働日"),
    (18, "年間削減 (USD/年・台)", f"=B17*{ELEC}", NUM2, "×電気単価"),
    (19, "参考：5台 (USD/年)", "=B18*5", NUM2, "原案No.1の5台"),
    (20, "調整案 対象台数（OL＋IL）", f"={I('1')}+{I('2')}", "#,##0",
     "新機がクラッチ機の稼働を置き換える前提（借用・レンタル機はクラッチ機）"),
    (21, "年間削減 合計 (USD/年)", "=B18*B20", NUM2, "13_費用対効果PLへ連携"),
]
for rr, lab, f, fmt, rem in srows:
    put(ws, f"A{rr}", lab, F_BOLD if rr in (18, 21) else F_BASE)
    if rr == 20:
        link(ws, f"B{rr}", f, fmt)
    else:
        calc(ws, f"B{rr}", f, fmt, bold=rr in (18, 21), fill=FILL_KEY if rr in (18, 21) else None)
    put(ws, f"D{rr}", rem)
section(ws, 23, "3. 品質・操作性（主効果）")
text_table(ws, 24, ["項目", "クラッチモーター（現状）", "サーボモーター（新）", "品質への影響"], [
    ("停止", "ペダルを離しても惰性回転で数針進む", "ペダルOFFで即停止", "曲線部で止めて向きを直す作業の精度が上がる"),
    ("曲線部の微調整", "微調整が難しく熟練に依存", "低速・針単位の調整が容易", "縫いズレ減少、品質安定"),
    ("属人性", "熟練者でないと品質がばらつく", "操作が標準化しやすい", "習熟期間短縮、不良低減"),
    ("糸切り", "手作業", "自動糸切り", "工程内の付帯作業削減"),
], heights=30)

# ================================================================ 07 省人効果（PS・Spec-D）
ws = sheet(S_PS, {"A": 6, "B": 10, "C": 22, "D": 7, "E": 11, "F": 7, "G": 8, "H": 7, "I": 8, "J": 7,
                  "K": 9, "L": 9, "M": 10, "N": 9, "O": 8, "P": 9, "Q": 9, "R": 12, "S": 12, "T": 9,
                  "U": 9, "V": 22, "W": 9, "X": 9, "Y": 9, "Z": 9})
title(ws, "省人効果：パターンシーマ（No.4・5）／Spec-D（No.7）",
      "省人化＝人員削減ではない。工程秒削減 → 必要工数削減 → 必要人員減 → 余力人員の再配置 → 工場全体の生産性向上（例：30人必要→28人で同じ生産量→2人を他工程へ）。")
put(ws, "A3", "1人年間有効稼働秒", F_BOLD, border=False)
link(ws, "E3", "=" + SECS, "#,##0")
note(ws, "F3", "年間削減工数（秒）＝（Before工程秒−After工程秒）×年間生産数量　／　省人工数＝年間削減工数÷1人年間有効稼働秒")
note(ws, "A4", "※全ケースでサイクルが短くなるわけではない（同等・数秒遅いケースあり）。サイクル基準と『作業者拘束時間』基準の両方を表示し、"
     "拘束時間・工程集約・品質・習熟で評価する。黄色セルに工場の実測raw dataを入力。")
section(ws, 5, "1. パターンシーマ：手作業 → 改善治具＋1本針 → パターンシーマ（同一工程の3方式比較）")
# 2段見出し
groups = [("A6:E6", "対象"), ("F6:G6", "手作業"), ("H6:I6", "治具＋1本針"), ("J6:N6", "パターンシーマ"),
          ("O6:U6", "差分・効果（手作業比）"), ("V6:V6", "判定"), ("W6:Z6", "定性効果（実測）")]
for rng, lab in groups:
    a = rng.split(":")[0]
    put(ws, a, lab, F_HEAD, PatternFill("solid", fgColor="2F75B5"), align=CENTER)
    if rng.split(":")[0] != rng.split(":")[1]:
        ws.merge_cells(rng)
header(ws, 7, ["No", "顧客", "部位・工程", "機種(大/中)", "年間生産数量(枚)", "工程数", "工程秒",
               "工程数", "工程秒", "工程数", "サイクル秒", "うち自動縫製秒", "自動縫製中の並行作業秒(実測)",
               "作業者拘束秒", "工程数削減", "サイクル差(秒)", "拘束時間差(秒)", "年間削減工数(秒)サイクル基準",
               "年間削減工数(秒)拘束基準", "省人工数(人)サイクル基準", "省人工数(人)拘束基準", "判定",
               "習熟日数 手作業", "習熟日数 PS", "不良率 手作業", "不良率 PS"])
ws.row_dimensions[7].height = 58
ws.freeze_panes = "D8"
dvm = DataValidation(type="list", formula1='"大,中"', allow_blank=True)
ws.add_data_validation(dvm)
EX, D0, D1 = 8, 9, 28


def ps_formulas(r):
    return {
        "N": f'=IF(K{r}="","",K{r}-N(M{r}))',
        "O": f'=IF(OR(F{r}="",J{r}=""),"",F{r}-J{r})',
        "P": f'=IF(OR(G{r}="",K{r}=""),"",G{r}-K{r})',
        "Q": f'=IF(OR(G{r}="",N{r}=""),"",G{r}-N{r})',
        "R": f'=IF(OR(E{r}="",P{r}=""),"",P{r}*E{r})',
        "S": f'=IF(OR(E{r}="",Q{r}=""),"",Q{r}*E{r})',
        "T": f'=IF(R{r}="","",R{r}/{SECS})',
        "U": f'=IF(S{r}="","",S{r}/{SECS})',
        "V": f'=IF(P{r}="","",IF(P{r}>0,"サイクル短縮",IF(P{r}=0,"サイクル同等：拘束・品質で評価","サイクル増：拘束・品質で評価")))',
    }


fmts = {"N": "#,##0", "O": "#,##0", "P": "#,##0;(#,##0);0", "Q": "#,##0;(#,##0);0", "R": USD, "S": USD,
        "T": NUM3, "U": NUM3, "V": "@"}
# 記入例（340/330/310秒・6/3/2工程は工場raw dataの例。数量・自動縫製秒・並行作業秒は仮置き）
ex = {"A": "記入例", "B": "例", "C": "ファスナー（例）", "D": "中", "E": 10000, "F": 6, "G": 340, "H": 3, "I": 330,
      "J": 2, "K": 310, "L": 180, "M": 120}
for col, v in ex.items():
    put(ws, f"{col}{EX}", v, F_EX, FILL_EX, fmt="#,##0" if isinstance(v, int) else None)
for col, f in ps_formulas(EX).items():
    put(ws, f"{col}{EX}", f, F_EX, FILL_EX, fmt=fmts[col])
for col in "WXYZ":
    put(ws, f"{col}{EX}", None, F_EX, FILL_EX)
ws[f"A{EX}"].comment = Comment("記入例：工程数と工程秒（340/330/310）は工場raw dataの例。年間数量・自動縫製秒・並行作業秒は仮置きのため合計に含めない。", "APIC")
for r in range(D0, D1 + 1):
    put(ws, f"A{r}", r - D0 + 1, align=CENTER)
    for col in "BCEFGHIJKLM":
        inp(ws, f"{col}{r}", None, "#,##0" if col not in "BC" else None)
    c = inp(ws, f"D{r}", None)
    dvm.add(c)
    for col, f in ps_formulas(r).items():
        calc(ws, f"{col}{r}", f, fmts[col])
    for col in "WX":
        inp(ws, f"{col}{r}", None, "#,##0")
    for col in "YZ":
        inp(ws, f"{col}{r}", None, PCT)
PT = D1 + 1
put(ws, f"A{PT}", "合計（記入例を除く）", F_BOLD, FILL_TOTAL)
ws.merge_cells(f"A{PT}:Q{PT}")
for col in "RSTU":
    calc(ws, f"{col}{PT}", f"=SUM({col}{D0}:{col}{D1})", fmts[col], bold=True, fill=FILL_TOTAL)
for col in "VWXYZ":
    put(ws, f"{col}{PT}", "", fill=FILL_TOTAL)
note(ws, f"A{PT+1}", "対象：TNF／GW／Montbell／HH の ファスナー・ポケット・フードキルト等。大型＝70×120（130×95機）、中型＝40×60。")

r0 = PT + 3
section(ws, r0, "2. シームシーリング Spec-D：ロングアーム工程（袖・裾など長尺パーツ）の2名体制 → 1名化")
header(ws, r0 + 1, ["No", "工程", "", "年間数量(枚)", "現行 工程秒/枚", "現行 人数", "Spec-D 工程秒/枚",
                    "Spec-D 人数", "現行 人秒/年", "", "Spec-D 人秒/年", "", "年間削減工数(秒)", "", "省人工数(人)", "", "備考"])
for a, b in (("B", "C"), ("I", "J"), ("K", "L"), ("M", "N"), ("O", "P"), ("Q", "U")):
    ws.merge_cells(f"{a}{r0+1}:{b}{r0+1}")
SEX, S0, S1 = r0 + 2, r0 + 3, r0 + 10


def sd_formulas(r):
    return {
        "I": f'=IF(OR(D{r}="",E{r}="",F{r}=""),"",E{r}*F{r}*D{r})',
        "K": f'=IF(OR(D{r}="",G{r}="",H{r}=""),"",G{r}*H{r}*D{r})',
        "M": f'=IF(OR(I{r}="",K{r}=""),"",I{r}-K{r})',
        "O": f'=IF(M{r}="","",M{r}/{SECS})',
    }


sdf = {"I": USD, "K": USD, "M": USD, "O": NUM3}
exs = {"A": "記入例", "B": "袖 ロングアーム（例）", "D": 10000, "E": 60, "F": 2, "G": 60, "H": 1}
for rr in range(SEX, S1 + 1):
    for a, b in (("B", "C"), ("I", "J"), ("K", "L"), ("M", "N"), ("O", "P"), ("Q", "U")):
        ws.merge_cells(f"{a}{rr}:{b}{rr}")
for col, v in exs.items():
    put(ws, f"{col}{SEX}", v, F_EX, FILL_EX, fmt="#,##0" if isinstance(v, int) else None)
for col, f in sd_formulas(SEX).items():
    put(ws, f"{col}{SEX}", f, F_EX, FILL_EX, fmt=sdf[col])
put(ws, f"Q{SEX}", "数値はすべて仮置き（形式の例）", F_EX, FILL_EX)
for rr in range(S0, S1 + 1):
    put(ws, f"A{rr}", rr - S0 + 1, align=CENTER)
    inp(ws, f"B{rr}", None)
    for col in "DEFGH":
        inp(ws, f"{col}{rr}", None, "#,##0")
    for col, f in sd_formulas(rr).items():
        calc(ws, f"{col}{rr}", f, sdf[col])
    inp(ws, f"Q{rr}", None)
ST = S1 + 1
put(ws, f"A{ST}", "合計（記入例を除く）", F_BOLD, FILL_TOTAL)
ws.merge_cells(f"A{ST}:L{ST}")
calc(ws, f"M{ST}", f"=SUM(M{S0}:N{S1})", USD, bold=True, fill=FILL_TOTAL)
ws.merge_cells(f"M{ST}:N{ST}")
calc(ws, f"O{ST}", f"=SUM(O{S0}:P{S1})", NUM3, bold=True, fill=FILL_TOTAL)
ws.merge_cells(f"O{ST}:P{ST}")

r1 = ST + 2
section(ws, r1, "3. 集計：省人工数と再配置可能人員（拘束時間基準）")
header(ws, r1 + 1, ["", "項目", "", "", "値", "", "", "", "備考"])
ws.merge_cells(f"B{r1+1}:D{r1+1}")
ws.merge_cells(f"E{r1+1}:H{r1+1}")
ws.merge_cells(f"I{r1+1}:U{r1+1}")
agg = [
    ("パターンシーマ 中型（No.4）省人工数", f'=SUMIF(D{D0}:D{D1},"中",U{D0}:U{D1})', NUM2, "07表1 D列＝中"),
    ("パターンシーマ 大型（No.5）省人工数", f'=SUMIF(D{D0}:D{D1},"大",U{D0}:U{D1})', NUM2, "07表1 D列＝大"),
    ("Spec-D（No.7）省人工数", f"=O{ST}", NUM2, "07表2"),
    ("省人工数 合計（人）", f"=SUM(E{r1+2}:E{r1+4})", NUM2, ""),
    ("再配置可能人員（人・整数）", f"=ROUNDDOWN(E{r1+5},0)", "#,##0", "削減ではなく他工程へ再配置"),
    ("金額換算（USD/年）", f'=IF({LABOR}="",0,E{r1+5}*{LABOR})', USD, "前提条件C14の人件費が未入力の間は0"),
]
PS_AGG = {}
for k, (lab, f, fmt, rem) in enumerate(agg):
    rr = r1 + 2 + k
    put(ws, f"B{rr}", lab, F_BOLD if k in (3, 4) else F_BASE)
    ws.merge_cells(f"B{rr}:D{rr}")
    calc(ws, f"E{rr}", f, fmt, bold=k in (3, 4), fill=FILL_KEY if k in (3, 4) else None)
    ws.merge_cells(f"E{rr}:H{rr}")
    put(ws, f"I{rr}", rem)
    ws.merge_cells(f"I{rr}:U{rr}")
    PS_AGG[k] = rr

# ================================================================ 08 PS運用設計
ws = sheet(S_OPS, {"A": 22, "B": 34, "C": 30, "D": 38, "E": 26})
title(ws, "パターンシーマ 運用設計（3/25要確認事項②）",
      "パターンシーマは導入することが目的にならないように、動線、人配置、縫製中の工員の動きを含めて運用設計する。")
section(ws, 3, "1. 原則")
text_table(ws, 4, ["原則", "内容"], [
    ("成果指標", "台数ではなく、工程集約・作業者拘束時間削減・属人性排除・習熟期間短縮・品質安定・必要工数削減で評価する"),
    ("APICで自動化が効く理由", "APICは多品種・切替型。同一品番を通年生産するなら人の熟練で最大効率に近づけるが、品番切替のたびに習熟がリセットされるため、"
     "熟練に依存しない自動機の優位性が大きい"),
    ("1人複数台", "前提にしない。動線とサイクルの実測で、自動縫製中に他方のセット・取り出しが収まることを確認できた場合のみ採用"),
    ("余力人員", "削減ではなく前後工程へ再配置（例：30人必要→28人で同じ生産量→2人を他工程へ）"),
], heights=36)
ws.merge_cells("B4:E4")
for rr in range(5, 9):
    ws.merge_cells(f"B{rr}:E{rr}")
section(ws, 10, "2. Before／After 作業フロー")
text_table(ws, 11, ["ステップ", "手作業（Before）", "改善治具＋1本針", "パターンシーマ（After）", "作業者の状態"], [
    ("① パーツ準備", "パーツを揃え、印付け", "パーツを揃える", "型枠（テンプレート）単位でパーツを揃える", "拘束"),
    ("② セット", "合わせ・ピン打ち", "治具にセット", "型枠にパーツをセット → 機械に装着", "拘束"),
    ("③ 縫製", "手縫い（例：6工程）", "1本針で縫製（例：3工程）", "自動縫製（例：2工程）", "Before：拘束／After：非拘束"),
    ("④ 自動縫製中", "－", "－", "2枚目の型枠に次製品をセット、前製品の糸処理・検品", "並行作業（実測対象）"),
    ("⑤ 取り出し・検品", "検品", "治具から外して検品", "型枠から取り出し・検品 → 次の型枠を装着", "拘束"),
], heights=32)
section(ws, 18, "3. 配置案")
text_table(ws, 19, ["機種", "配置", "理由", "留意点", "確認者"], [
    ("中型 60×40（No.4）", "対象ラインのポケットファスナー工程位置にライン内配置", "機体が小さくライン内に入る（2/25確認）",
     "型枠2枚運用で縫製中に次をセット。前後工程との仕掛り置場", "ライン長・IE"),
    ("大型 130×95（No.5）", "複数ライン共用セル（フロントファスナー・キルティング）", "品番により負荷が大きく変動",
     "搬送距離、仕掛り置場、ライン間の優先順位ルール", "生産管理・IE"),
    ("型枠（テンプレート）", "品番別に製作・保管・識別", "多品種切替では型枠準備が律速になる",
     "新品番の投入前に型枠製作リードタイムを確保（生産準備ワークフローへ組込み）", "技術"),
], heights=40)
section(ws, 24, "4. 人配置の判断条件（1人2台にする場合）")
text_table(ws, 25, ["条件", "判定方法"], [
    ("自動縫製秒 ≥ もう1台のセット秒＋取り出し秒＋検品秒＋移動秒", "07表1の実測値（自動縫製秒・並行作業秒）で確認"),
    ("2台間の移動距離が短い（目安：数歩以内）", "レイアウト図で確認"),
    ("待ち（機械待ち・人待ち）が発生しない", "10サイクル以上のビデオ測定で確認"),
], heights=28)
ws.merge_cells("B25:E25")
for rr in range(26, 29):
    ws.merge_cells(f"B{rr}:E{rr}")
section(ws, 30, "5. 測定計画")
text_table(ws, 31, ["測定項目", "方法", "回数", "時期", "担当"], [
    ("工程秒（Before／After）", "ビデオ撮影＋ストップウォッチ", "各工程10サイクル以上", "導入前・導入1か月後", "IE・工程担当"),
    ("作業者拘束時間・並行作業時間", "自動縫製中の作業内容を記録", "同上", "同上", "IE"),
    ("習熟期間", "新人が標準時間に到達するまでの日数", "対象者ごと", "導入後", "ライン長"),
    ("不良率", "工程内不良・最終検品", "1か月集計", "導入前後", "QC"),
    ("稼働率", "PS稼働時間÷就業時間", "週次", "導入後", "生産管理"),
], heights=28)

# ================================================================ 09 スポットプレス
ws = sheet(S_SPOT, {"A": 44, "B": 14, "C": 14, "D": 14, "E": 30, "F": 16, "G": 36})
title(ws, "スポットプレス 必要台数と入替計画（No.11・最優先／安全投資）",
      "安全装置なしの旧型は安全上の理由で全数入替。ROIで判断しない。必要総数＝対象ライン数×必要台数/ライン＋共用分＋予備・保全分。")
section(ws, 3, "A. 現有機の状況")
header(ws, 4, ["区分", "台数", "年式・使用年数", "安全装置", "後付け可否", "修理可否", "備考"])
put(ws, "A5", "旧型（アーム斜め動作）")
inp(ws, "B5", None, "#,##0", comment="旧型の現有台数を入力")
put(ws, "C5", "約15年")
put(ws, "D5", "無", align=CENTER)
put(ws, "E5", "不可（メーカー確認、廃棄推奨）")
inp(ws, "F5", None)
put(ws, "G5", "圧力計算不可・3L等に使用不可。圧力ムラで再プレス増")
put(ws, "A6", "新型 PTTG 30×10cm")
inp(ws, "B6", 15, "#,##0")
inp(ws, "C6", None)
put(ws, "D6", "有", align=CENTER)
put(ws, "E6", "－")
put(ws, "F6", "－")
put(ws, "G6", "原案S14・R14")
section(ws, 8, "B. 必要台数の計算")
header(ws, 9, ["項目", "", "値", "", "備考", "", ""])
for rr in range(9, 28):
    ws.merge_cells(f"A{rr}:B{rr}")
    ws.merge_cells(f"E{rr}:G{rr}")
spot = [
    (10, "対象ライン数", None, "inp", "#,##0", "スポットプレスを使う量産ライン数"),
    (11, "1ライン当たり必要台数", None, "inp", "0.0", "ライン別使用状況（C表）から設定"),
    (12, "ライン計", '=IF(OR(C10="",C11=""),"",C10*C11)', "calc", "#,##0", ""),
    (13, "共用分（仕上げ・サンプル・部分プレス等）", None, "inp", "#,##0", ""),
    (14, "予備・保全分", None, "inp", "#,##0", "故障・修理中の代替"),
    (15, "計算上の必要総数", '=IF(C12="","",ROUNDUP(C12+N(C13)+N(C14),0))', "calc", "#,##0", ""),
    (16, "暫定 必要総数（3/25概算45〜50台の下限）", 45, "inp", "#,##0", "計算根拠の入力が終わるまでの暫定値"),
    (17, "採用 必要総数", "=IF(C15=\"\",C16,C15)", "calc", "#,##0",
     '=IF(C15="","※暫定値を使用中：上の計算欄を入力すると自動で置き換わる","計算値を使用")'),
    (18, "新型（安全装置付）現有", "=B6", "calc", "#,##0", ""),
    (19, "原案① 計上台数", 10, "inp", "#,##0", ""),
    (20, "2026年度 購入必要数", "=MAX(0,C17-C18)", "calc", "#,##0", "→ 02_投資総括 No.11 の調整案台数"),
    (21, "うち 更新分（旧型の置換）", '=IF(B5="","旧型台数入力待ち",MIN(C20,B5))', "calc", "#,##0", ""),
    (22, "うち 純増分", '=IF(B5="","旧型台数入力待ち",C20-C21)', "calc", "#,##0", ""),
    (23, "原案との差（台）", "=C20-C19", "calc", "#,##0;(#,##0);0", ""),
    (24, "単価（USD）", "=" + ref(S_SUM, f"$F${SUM_ROW['11']}"), "link", USD, "02_投資総括"),
    (25, "購入金額（USD）", "=C20*C24", "calc", USD, ""),
    (26, "購入金額（円）", f"=C25*{RATE}", "calc", YEN, ""),
    (27, "購入後の旧型（安全装置なし）", "全数使用停止・廃棄", "text", None, "必要総数をすべて安全装置付で充足"),
]
for rr, lab, v, kind, fmt, rem in spot:
    put(ws, f"A{rr}", lab, F_BOLD if rr in (17, 20) else F_BASE)
    if kind == "inp":
        inp(ws, f"C{rr}", v, fmt)
    elif kind == "link":
        link(ws, f"C{rr}", v, fmt)
    elif kind == "calc":
        calc(ws, f"C{rr}", v, fmt, bold=rr in (17, 20), fill=FILL_KEY if rr in (17, 20) else None)
    else:
        put(ws, f"C{rr}", v)
    ws.merge_cells(f"C{rr}:D{rr}")
    if isinstance(rem, str) and rem.startswith("="):
        calc(ws, f"E{rr}", rem)
    else:
        put(ws, f"E{rr}", rem)
section(ws, 29, "C. ライン別 使用状況（現況調査）")
header(ws, 30, ["ライン・工程", "旧型 使用台数", "新型 使用台数", "ピーク時 必要台数", "主な用途", "", "備考"])
for rr in range(31, 43):
    for col in "ABCDEG":
        inp(ws, f"{col}{rr}", None, "#,##0" if col in "BCD" else None)
    ws.merge_cells(f"E{rr}:F{rr}")
put(ws, "A43", "合計", F_BOLD, FILL_TOTAL)
for col in "BCD":
    calc(ws, f"{col}43", f"=SUM({col}31:{col}42)", "#,##0", bold=True, fill=FILL_TOTAL)
section(ws, 45, "D. 証跡チェックリスト")
text_table(ws, 46, ["証跡", "内容", "状態"], [
    ("旧型の写真", "安全装置がないことが分かる写真（挟まれ箇所）", ""),
    ("メーカー見解", "安全装置の後付け不可・廃棄推奨の書面またはメール", ""),
    ("年式・台帳", "固定資産台帳の取得年・簿価", ""),
    ("故障・修理履歴", "修理可否、部品入手可否", ""),
    ("ヒヤリハット", "挟まれ・やけど等の記録（あれば）", ""),
    ("品質影響", "圧力ムラによる再プレス回数、耐水圧不良との関係", ""),
])
for rr in range(47, 53):
    inp(ws, f"C{rr}", None)
section(ws, 54, "E. 説明文案")
put(ws, "A55", "旧型スポットプレスは約15年経過し、安全装置の後付けができずメーカーも廃棄を推奨している。安全上の問題があるため、"
    "費用対効果ではなく安全確保を理由に最優先で全数入替を行う。必要総数は［対象ライン数×必要台数/ライン＋共用分＋予備分］で算出し、"
    "安全装置付の新型現有15台との差を2026年度に購入する。", border=False)
ws.merge_cells("A55:G57")
ws.row_dimensions[55].height = 30

# ================================================================ 10 平プレス検証
ws = sheet(S_FLAT, {"A": 16, "B": 46, "C": 36, "D": 22, "E": 14})
title(ws, "No.10 平プレス（125×83cm） プレス精度の検証方法（3/25要確認事項⑤）",
      "大型だからこそ圧力・温度の均一性が懸念にならないよう、受入時と定期で確認する。判定基準の数値はメーカー（H&H／JUKI）規格を取得して記入。")
text_table(ws, 3, ["検証項目", "検証方法", "判定基準", "確認頻度", "担当"], [
    ("水平度", "水平器で熱板の4隅を測定（必要に応じ中央を追加）", "4点とも許容範囲内（許容値：メーカー規格）", "受入時・月次・移設/修理後", "保全"),
    ("圧着均一性", "紙挟みテスト：熱板全面に紙片を配置（例：3×3＝9点）してプレスし、圧痕・引抜き抵抗の均一性を確認",
     "全点で圧痕が均一、抜け・弱い点がない", "受入時・週次・品番切替時", "技術"),
    ("圧力分布（補足）", "感圧紙（圧力測定フィルム）で熱板全面の圧力分布を可視化", "発色ムラが許容範囲内", "受入時・年次・異常時", "技術／メーカー"),
    ("温度", "接触式表面温度計で熱板9点を測定", "設定温度±許容差内（メーカー規格）", "始業時（中央1点）・月次（9点）", "工程リーダー／保全"),
    ("圧力", "圧力計の表示と設定値を照合", "設定値±許容差内（メーカー規格）", "始業時", "工程リーダー"),
    ("接着品質", "実素材（PUR接着・TNFエアチャンバー）のテストピースで剥離強度・耐水圧・外観を確認", "顧客／社内規格を満たす",
     "受入時・新品番立上げ時", "品質"),
    ("メーカー校正・保守", "メーカーによる温度・圧力のキャリブレーションと定期保守", "校正記録を保管", "年1回（頻度はメーカー推奨を確認）", "保全／メーカー"),
], heights=40)
note(ws, "A12", "記録：検証結果は日付・測定値・判定・担当を記録表に残し、異常時はプレス停止→保全連絡→再検証の手順とする。")

# ================================================================ 11 縮絨
ws = sheet(S_FELT, {"A": 6, "B": 22, "C": 46, "D": 34, "E": 16, "F": 12})
title(ws, "縮絨機 検討ステップ（3/25要確認事項⑦）",
      "結論：2026年度予算には計上しない。QN事例確認と小規模試験の結果で導入／見送りを判断する（試験後に導入判断）。")
text_table(ws, 3, ["No", "ステップ", "内容", "判断基準", "担当", "状態"], [
    (1, "QN事例確認", "QNで使用している生地、条件（温度・時間・張力）、改善した品質項目", "事例と効果データがあるか", "TIVN・加技部", ""),
    (2, "対象生地の選定", "APICで寸法・風合いが不安定な生地（サプライヤー別）", "不良・クレーム実績", "APIC品質", ""),
    (3, "使用条件", "どの工程で使うか（裁断前の生地安定化など）", "工程・LTへの影響", "APIC技術", ""),
    (4, "品質改善効果", "寸法変化率・斜行・外観の Before／After", "改善幅が規格内に収まるか", "APIC品質", ""),
    (5, "APIC適用性", "設置スペース、処理能力、リードタイム影響、外注の可否", "内製の必要性", "APIC生産管理", ""),
    (6, "小規模試験", "QNまたは外部で小ロット試験", "4の基準を満たすか", "TIVN", ""),
    (7, "導入／見送り判断", "効果が数値で確認できた場合のみ次年度以降に予算化", "－", "TIVN・TIJ", ""),
], heights=30)
for rr in range(4, 11):
    inp(ws, f"F{rr}", None)
note(ws, "A12", "経緯：過去にAPICでウールの縮絨テストを実施したが効果は明確でなかった。生地が不安定なサプライヤーが多く、工場側で生地を安定させる手段として検討。")

# ================================================================ 12 レンタル費
ws = sheet(S_RENT, {"A": 8, "B": 30, "C": 8, "D": 12, "E": 12, "F": 12, "G": 10, "H": 13, "I": 60})
title(ws, "外部レンタル費と投資による回避可能額（3/25要確認事項①）",
      "投資でレンタルが全てなくなるわけではない。回避率（G列）を機種別に設定し、回避可能な部分のみを効果として計上する。")
put(ws, "B3", "2025年 レンタル費 総額（USD）", F_BOLD)
inp(ws, "D3", 4890.75, USD2)
ws.merge_cells("D3:E3")
put(ws, "F3", "出所：APIC_2025年 機械レンタル費用総額.xlsx", F_NOTE, border=False, align=Alignment(vertical="top"))
header(ws, 5, ["年", "区分・機種", "台数", "期間", "金額(USD)", "対応投資キー", "回避率", "回避可能額(USD)", "出所・備考"])
rent = [
    ("2025", "大型パターンシーマ", None, None, 3016.54, "PS大", 1.0,
     "回避率100%は『No.5購入台数で繁忙期需要を賄える』前提（仮）。レンタル台数・期間を確認"),
    ("2025", "その他機種（内訳未確認）", None, None, "=D3-E6", "", 0,
     "内訳確認後に機種別へ分割。原案No.6の13,440USD/年はこの金額（最大1,874.21USD）と整合しない"),
    ("2024", "ロックミシン（10〜12月、約10台）", 10, "3か月", None, "ロック", None, "金額はD/Nから入力。回避率は購入後も残る不足（05）を考慮して設定"),
    ("2024", "その他", None, None, None, "", None, ""),
    ("2026", "実績（年初〜）", None, None, None, "", None, ""),
]
for k, (yr, kind, qty, per, amt, key, rate, rem) in enumerate(rent):
    rr = 6 + k
    put(ws, f"A{rr}", yr, align=CENTER)
    put(ws, f"B{rr}", kind)
    inp(ws, f"C{rr}", qty, "#,##0")
    inp(ws, f"D{rr}", per)
    if isinstance(amt, str):
        calc(ws, f"E{rr}", amt, USD2)
    else:
        inp(ws, f"E{rr}", amt, USD2)
    inp(ws, f"F{rr}", key)
    inp(ws, f"G{rr}", rate, PCT)
    calc(ws, f"H{rr}", f'=IF(OR(E{rr}="",G{rr}=""),"",E{rr}*G{rr})', USD2)
    put(ws, f"I{rr}", rem)
    ws.row_dimensions[rr].height = 30
for rr in range(11, 15):
    for col in "ABCDEFGI":
        inp(ws, f"{col}{rr}", None, {"C": "#,##0", "E": USD2, "G": PCT}.get(col))
    calc(ws, f"H{rr}", f'=IF(OR(E{rr}="",G{rr}=""),"",E{rr}*G{rr})', USD2)
put(ws, "A15", "合計", F_BOLD, FILL_TOTAL)
ws.merge_cells("A15:D15")
calc(ws, "E15", "=SUM(E6:E14)", USD2, bold=True, fill=FILL_TOTAL)
put(ws, "F15", "", fill=FILL_TOTAL)
put(ws, "G15", "", fill=FILL_TOTAL)
calc(ws, "H15", "=SUM(H6:H14)", USD2, bold=True, fill=FILL_TOTAL)
put(ws, "I15", "年度が混在（直近実績ベース）。同一機種の重複計上に注意", fill=FILL_TOTAL)
note(ws, "A16", "対応投資キー：ロック／PS大／PS中／DUMA／SpecD／温冷／平プレス（13_費用対効果PLが同じキーで集計）。")
section(ws, 18, "明細（D/N・Invoice単位）")
header(ws, 19, ["年月", "D/N・Invoice No.", "台数", "期間(日)", "金額(USD)", "機種・型式", "style", "line", "回避可否・理由"])
for rr in range(20, 45):
    for col in "ABCDEFGHI":
        inp(ws, f"{col}{rr}", None, {"C": "#,##0", "D": "#,##0", "E": USD2}.get(col))
put(ws, "A45", "合計", F_BOLD, FILL_TOTAL)
ws.merge_cells("A45:D45")
calc(ws, "E45", "=SUM(E20:E44)", USD2, bold=True, fill=FILL_TOTAL)

# ================================================================ 13 費用対効果・P/L
ws = sheet(S_ROI, {"A": 26, "B": 8, "C": 12, "D": 12, "E": 12, "F": 10, "G": 12, "H": 12, "I": 11, "J": 12,
                   "K": 12, "L": 12, "M": 50})
title(ws, "費用対効果（改善投資）と当年度P/L影響",
      "年間純効果＝レンタル削減＋外注費削減＋省人効果＋品質効果＋電力削減−追加保守・消耗品　／　単純回収年数＝投資額÷年間純効果")
header(ws, 4, ["投資案件", "原案No.", "投資額(USD)", "レンタル削減", "外注費削減", "省人工数(人)", "省人効果",
               "品質効果", "電力削減", "追加保守・消耗品", "年間純効果(USD/年)", "単純回収年数", "評価・備考"])
RENT_KEY = ref(S_RENT, "$F$6:$F$14")
RENT_AMT = ref(S_RENT, "$H$6:$H$14")
roi = [
    ("ロックミシン OL＋IL", "1,2", f"={J('1')}+{J('2')}", "ロック", None, None, "=" + ref(S_SERVO, "$B$21"),
     "主目的は不足解消（生産継続）。電力は副次効果"),
    ("パターンシーマ 中型", "4", f"={J('4')}", "PS中", None, "=" + ref(S_PS, f"$E${PS_AGG[0]}"), None,
     "省人は07の実測入力後に確定"),
    ("パターンシーマ 大型", "5", f"={J('5')}", "PS大", 3000, "=" + ref(S_PS, f"$E${PS_AGG[1]}"), None,
     "外注費3,000USD＝25FWキルティング外注（原案R8）。年間額（SS＋FW）を確認"),
    ("本縫針送りミシン DUMA", "6", f"={J('6')}", "DUMA", None, None, None, "レンタル削減は2025明細で確認（原案の13,440USD/年は不整合）"),
    ("シームシーリング Spec-D", "7", f"={J('7')}", "SpecD", None, "=" + ref(S_PS, f"$E${PS_AGG[2]}"), None,
     "2名体制→1名化の効果"),
    ("温冷プレス 60×12", "9", f"={J('9')}", "温冷", None, None, None, "工程秒約半減（原案）→ 削減秒×数量で省人を算出し入力"),
    ("平プレス 125×83", "10", f"={J('10')}", "平プレス", None, None, None, "プレス回数削減を品番別に計測"),
]
R0 = 5
for k, (lab, no, inv, key, outs, ppl, elec, rem) in enumerate(roi):
    rr = R0 + k
    put(ws, f"A{rr}", lab)
    put(ws, f"B{rr}", no, align=CENTER)
    link(ws, f"C{rr}", inv, USD)
    link(ws, f"D{rr}", f'=SUMIF({RENT_KEY},"{key}",{RENT_AMT})', USD2)
    inp(ws, f"E{rr}", outs, USD2)
    if ppl:
        link(ws, f"F{rr}", ppl, NUM2)
        calc(ws, f"G{rr}", f'=IF({LABOR}="",0,F{rr}*{LABOR})', USD2)
    else:
        inp(ws, f"F{rr}", None, NUM2)
        calc(ws, f"G{rr}", f'=IF(OR({LABOR}="",F{rr}=""),0,F{rr}*{LABOR})', USD2)
    inp(ws, f"H{rr}", None, USD2)
    if elec:
        link(ws, f"I{rr}", elec, USD2)
    else:
        inp(ws, f"I{rr}", None, USD2)
    inp(ws, f"J{rr}", None, USD2)
    calc(ws, f"K{rr}", f"=N(D{rr})+N(E{rr})+N(G{rr})+N(H{rr})+N(I{rr})-N(J{rr})", USD2, bold=True)
    if k == 0:  # ロックは電力だけで回収年数を出さない（主効果はレンタル回避）
        calc(ws, f"L{rr}", f'=IF(D{rr}=0,"レンタル実績額 入力待ち",C{rr}/K{rr})', '0.0"年"')
    else:
        calc(ws, f"L{rr}", f'=IF(K{rr}>0,C{rr}/K{rr},"効果データ待ち")', '0.0"年"')
    put(ws, f"M{rr}", rem)
    ws.row_dimensions[rr].height = 30
RT = R0 + len(roi)
put(ws, f"A{RT}", "改善投資 合計", F_BOLD, FILL_TOTAL)
put(ws, f"B{RT}", "", fill=FILL_TOTAL)
for col in "CDEFGHIJK":
    calc(ws, f"{col}{RT}", f"=SUM({col}{R0}:{col}{RT-1})", NUM2 if col == "F" else USD2, bold=True, fill=FILL_TOTAL)
calc(ws, f"L{RT}", f'=IF(COUNTIF(L{R0}:L{RT-1},"*待ち")>0,"一部データ待ち",C{RT}/K{RT})', '0.0"年"', bold=True, fill=FILL_TOTAL)
put(ws, f"M{RT}", "", fill=FILL_TOTAL)
note(ws, f"A{RT+1}", "省人効果（G列）は前提条件C14（人件費）入力後に金額化。省人は再配置であり、人件費の削減ではなく『追加採用の回避・生産量増』として説明する。")

r = RT + 3
section(ws, r, "ROI対象外（安全・顧客要求・検証）")
text_table(ws, r + 1, ["投資案件", "原案No.", "投資額(USD)", "判断理由"], [
    ("スポットプレス", "11", None, "安全（安全装置なし旧型の全数入替）。金額換算しない"),
    ("Nawon温冷プレス 4種", "12-15", None, "顧客要求（26FW GW要請）。受注維持"),
    ("Spec-IJ・US-512", "8,18", None, "戦略／検証（少数導入して評価）"),
])
link(ws, f"C{r+2}", f"={J('11')}", USD)
link(ws, f"C{r+3}", f"={J('12')}+{J('13')}+{J('14')}+{J('15')}", USD)
link(ws, f"C{r+4}", f"={J('8')}+{J('18')}", USD)
for rr in range(r + 1, r + 5):
    ws.merge_cells(f"D{rr}:M{rr}")

r = r + 6
section(ws, r, "当年度P/L影響（年額・円）")
note(ws, f"A{r+1}", "当年度P/L影響＝新規設備年間償却−既存設備償却減少額−年間削減効果　（＋＝費用増 → 安全・顧客要求・生産継続の必要性で説明／−＝P/L改善）")
header(ws, r + 2, ["項目", "", "ケースA", "ケースB", "", "備考"])
ws.merge_cells(f"A{r+2}:B{r+2}")
ws.merge_cells(f"F{r+2}:M{r+2}")
pl = [
    ("調整案 投資総額（USD）", f"={ref(S_SUM, f'$J${SUM_T}')}", f"={ref(S_SUM, f'$J${SUM_T}')}", USD, "02_投資総括"),
    ("うちTI償却対象（USD）", f"=C{r+3}", f"=C{r+3}-{APIC}", USD,
     "A：APIC負担分もTI資産として償却（保守的）／B：APIC負担50,000USDをTI償却対象から除外。資産保有者・償却処理を確認するまでAを基本とする"),
    ("新規設備 年間償却（円）", f"=C{r+4}*{RATE}/{YEARS}", f"=D{r+4}*{RATE}/{YEARS}", YEN, "定額・償却年数は前提条件"),
    ("既存設備 償却減少額（円）", None, f"=C{r+6}", YEN,
     "入替で除却する既存機の年間償却。ロック旧機・旧型スポットプレスは償却済が多い（例：原案V4＝W4）→ 0なら空欄"),
    ("年間削減効果（円）", f"=K{RT}*{RATE}", f"=C{r+7}", YEN, "改善投資の年間純効果（上表）"),
    ("当年度P/L影響（円）", f"=C{r+5}-N(C{r+6})-C{r+7}", f"=D{r+5}-N(D{r+6})-D{r+7}", YEN, ""),
    ("判定", f'=IF(C{r+8}>0,"費用増：安全・顧客要求・生産継続で説明","P/L改善")',
     f'=IF(D{r+8}>0,"費用増：安全・顧客要求・生産継続で説明","P/L改善")', None, ""),
]
for k, (lab, a, b, fmt, rem) in enumerate(pl):
    rr = r + 3 + k
    put(ws, f"A{rr}", lab, F_BOLD if k in (5, 6) else F_BASE)
    ws.merge_cells(f"A{rr}:B{rr}")
    if a is None:
        inp(ws, f"C{rr}", None, fmt)
    elif k == 0:
        link(ws, f"C{rr}", a, fmt)
    else:
        calc(ws, f"C{rr}", a, fmt, bold=k in (5, 6), fill=FILL_KEY if k in (5, 6) else None)
    if k == 0:
        link(ws, f"D{rr}", b, fmt)
    else:
        calc(ws, f"D{rr}", b, fmt, bold=k in (5, 6), fill=FILL_KEY if k in (5, 6) else None)
    put(ws, f"E{rr}", "")
    put(ws, f"F{rr}", rem)
    ws.merge_cells(f"F{rr}:M{rr}")
    ws.row_dimensions[rr].height = 30 if len(rem) > 40 else 18

# ================================================================ 14 原案チェック
ws = sheet(S_CHK, {"A": 5, "B": 22, "C": 54, "D": 40, "E": 44})
title(ws, "原案（20260403版）の整合性チェック ― 提出前の修正候補",
      "TIJから指摘される前に修正しておくべき箇所。原案ファイル『2026年度 設備投資案』シートのセル番地で記載。")
text_table(ws, 3, ["No", "箇所", "内容", "影響", "対応案"], [
    (1, "R9（No.6 DUMA）", "年間レンタル費削減額 112USD×12か月×10台＝13,440USD/年、回収0.7年と記載",
     "2025年のレンタル費総額は4,890.75USD（大型PSを除くと1,874.21USD）で、記載の削減額は実績と整合しない。数値根拠の信頼性が問われる",
     "2025明細の針送りミシン分に置き換える。または『レンタル単価112USD/月・台』を前提とした想定値と明記する"),
    (2, "R4（No.1）", "『1ラインあたり平均4台稼働』と記載", "31ライン×4台＝124台となり必要台数が過大に見える",
     "製品構成（80%×3台＋20%×4台）＝平均3.2台、通常100台／ピーク112台に統一（05）"),
    (3, "R4・S4（No.1）", "『C判定機種が4台』", "新ロジック（老朽化重大故障4台・修理復旧不可・実稼働対象外）と表現が異なる",
     "『修理復旧不可4台』に統一"),
    (4, "R5・S5（No.2）", "R5『現状67台が稼働』、S5『68台』", "数値不一致", "現有68台（うちサンプル固定4台、生産用64台）に統一"),
    (5, "I25・J25・O25", "合計式が =SUM(I12:I24) 等で4〜11行目を含まない",
     "優先②の台数が10台（実際38台）、②金額が258,730USD（実際355,230USD）と過小表示", "合計式を SUM(I4:I24) 等に修正"),
    (6, "M27", "SUM式の参照先が削除されている（参照切れ）", "エラー値が表示される", "削除または参照を修正"),
    (7, "A21・A22", "No.18が2行重複", "以降の番号がずれる", "連番を振り直す"),
    (8, "N6・N16・N19・N22〜N24", "優先①金額（N列）の式が優先②台数（I列）×L列を参照", "現在は0表示で実害はないが式が不統一",
     "N列の式を『H列×L列』に統一"),
    (9, "X18", "残存簿価が『-』（V18・W18は入力済）", "簿価が表示されない", "X18をV18−W18の式に修正"),
    (10, "O26・P26・P30", "為替150円と160円が混在", "資料により円換算額が異なる", "1レートに統一（本資料は160円）"),
    (11, "R8（No.5）", "キルティング外注『25FWで約3千USD』", "年間額か季節額か不明", "SS＋FWの年間額を確認"),
    (12, "議事メモシート", "発言者の個人名を含む発言録が同じブックにある", "公式報告書に個人名を記載しない方針",
     "提出版から議事メモシートを外すか、発言者名を削除"),
    (13, "3.27確認事項シート", "日付が2025/3/27（前年度の説明会資料）", "2026年度資料と混同される", "提出版から外す"),
], heights=48)

# ================================================================ 03 要確認事項
ws = sheet(S_QA, {"A": 6, "B": 30, "C": 50, "D": 40, "E": 18, "F": 36, "G": 16})
title(ws, "要確認事項への回答（3/25打合せ ①〜⑦＋新規条件）",
      "現状 → 課題 → 数値根拠 → 導入効果 → 費用対効果 → 投資妥当性 → 優先順位 の順に説明する。")
qa = [
    ("①", "機材導入で削減できる年間外部レンタル費。パターンシーマ・Spec-D導入による要員減の数値",
     "レンタル費は機種別に『投資で回避できる部分』だけを効果計上（100%計上しない）。省人化は人員削減ではなく『工程秒削減→必要工数削減→必要人員減→余力人員の再配置→工場全体の生産性向上』と定義し、Before/After実測で算出。",
     "2025年レンタル総額4,890.75USD（うち大型PS 3,016.54USD）。省人工数＝年間削減工数÷1人年間有効稼働秒（8h×295日）。",
     f"{S_RENT}\n{S_PS}\n{S_ROI}",
     "2024・2026年レンタル明細（機種・台数・期間・D/N・style/line）、PS・Spec-Dの実測（工程秒・拘束秒）、対象品番の年間数量", "枠組み完成・データ入力待ち"),
    ("②", "パターンシーマ導入時の動線・工員配置・縫製中の動きの設計",
     "導入を目的化しない。中型はライン内配置、大型は共用セル。自動縫製中に次製品をセットする型枠2枚運用を前提とし、1人複数台は実測で並行作業時間を確認できた場合のみ採用。",
     "Before/Afterフロー（作業者→セット→自動縫製→縫製中に次製品準備→取り出し）と測定計画（各工程10サイクル以上）", S_OPS,
     "現場レイアウト図、型枠製作体制、実測結果", "運用案作成済・現場検証待ち"),
    ("③", "スポットプレスを最優先で入替。必要台数の計算根拠と現在の運用状況",
     "安全装置なしの旧型（約15年、後付け不可・メーカー廃棄推奨）は安全上の理由で全数入替。必要総数＝対象ライン数×必要台数/ライン＋共用分＋予備・保全分。ROIではなく安全で判断。更新分と純増分を分けて表示。",
     f'="暫定：必要"&TEXT({ref(S_SPOT, "$C$17")},"0")&"台−新型現有"&TEXT({ref(S_SPOT, "$C$18")},"0")&"台＝購入"&TEXT({ref(S_SPOT, "$C$20")},"0")&"台（原案10台）"',
     S_SPOT, "旧型台数・年式・修理可否、ライン別使用台数、ピーク需要、写真・ヒヤリハット", "計算式完成・現況データ待ち"),
    ("④", "ダウン充填機の優先度を下げる", "2026年度は購入見送り（現有4台で最大10ライン対応可）", "－", S_SUM, "－", "対応済"),
    ("⑤", "No.10平プレスのプレス精度の検証方法",
     "水平器4点測定、紙挟みテスト（＋感圧紙）、温度・圧力の実測、実素材での接着評価、メーカー校正を組み合わせ、受入時と定期で確認。",
     "検証項目｜検証方法｜判定基準｜確認頻度｜担当 の表で定義", S_FLAT, "判定基準のメーカー規格値、校正周期", "案作成済・メーカー確認待ち"),
    ("⑥", "定常投資とは別に48百万円枠を利用できるか", "本資料の対象外（TIJ側の確認事項）。本案は定常投資枠内・4,000万円以下で構成。", "－", "－", "－", "TIJ確認"),
    ("⑦", "縮絨機の導入検討（QN事例を含め確認）",
     "2026年度は購入せず、QN事例確認→対象生地→使用条件→品質改善効果→APIC適用性→小規模試験→導入／見送り判断。",
     "APIC過去試験（ウール）は効果が明確でなかった", S_FELT, "QN事例（生地・条件・効果）", "試験計画作成済"),
    ("新", "APICが投資額のうち50,000USDを負担",
     "負担対象の機種・資産保有者・償却処理を確認するまで、TI償却の減少として自動計上しない（P/LはケースA/Bを併記）。",
     f'="TI負担＝調整案"&TEXT({ref(S_SUM, f"$J${SUM_T}")},"#,##0")&"USD−50,000USD"',
     f"{S_SUM}\n{S_ROI}", "負担対象機種、資産保有者、償却処理", "要確認"),
    ("新", "投資額はSuntoolを除き概ね4,000万円以下",
     "安全・顧客要求・生産継続を確保し、改善投資は段階導入（効果確認後に追加）として調整。",
     f'="調整案 "&TEXT({ref(S_SUM, f"$K${SUM_T}")}/1000000,"0.00")&"百万円（原案① "&TEXT({ref(S_SUM, f"$H${SUM_T}")}*{RATE}/1000000,"0.00")&"百万円）"',
     S_SUM, "為替レートの統一（150円／160円）", "調整案作成済"),
]
header(ws, 3, ["No", "確認事項", "回答骨子", "数値・根拠", "根拠シート", "未確定データ・次アクション", "状態"])
for k, row in enumerate(qa):
    rr = 4 + k
    for i, v in enumerate(row):
        col = "ABCDEFG"[i]
        if isinstance(v, str) and v.startswith("="):
            calc(ws, f"{col}{rr}", v)
        else:
            put(ws, f"{col}{rr}", v, align=CENTER if col in "AG" else WRAP)
    ws.row_dimensions[rr].height = 80

# ================================================================ 04 機種別比較
ws = sheet(S_CMP, {"A": 6, "B": 7, "C": 20, "D": 13, "E": 36, "F": 32, "G": 22, "H": 30, "I": 32, "J": 24, "K": 14})
title(ws, "機種別比較（現状方式｜改善・代替案｜新設備案｜差分｜年間効果｜投資判断）",
      "ダウン充填装置比較の形式を踏襲。購入以外の代替案を必ず比較し、効果は数値で確認できるものだけを記載。")
cmp_rows = [
    (1, "11", "スポットプレス", "必須（安全）",
     "旧型：約15年経過、アームが斜め動作で圧力計算不可、安全装置なし。3L等の厳しい条件に使えず、圧力ムラで再プレス増。新型30×10cm 15台保有",
     "安全装置の後付け：不可（メーカー確認、廃棄推奨）／修理継続：安全リスクが残る", "PTTG 30×10cm（安全装置付）764USD/台",
     "安全装置あり、垂直加圧で圧力管理可、3L対応", "安全リスク解消（金額換算しない）。再プレス回数は導入後に計測", "最優先。台数＝09の計算（暫定30台）", S_SPOT),
    (2, "12-15", "Nawon温冷プレス 4種", "必須（顧客要求）",
     "現状のシーム機だけでは対応できない品番あり（26FW GW要請、SW1課から4種の導入要望）。2025年度にHTM-8601を1台導入済",
     "H&H同等機：CSL-206 9,500／CSL-213 14,500／CSL-207 10,500USD。カーブ部はNawonがアーム調整の自由度で優位、ストレート型は性能差が小さい",
     "Nawon HTM-8749×2、8739×1、8709×1、8601×1", "フロントジッパー・袖・裾・フード曲線を一度で均一にプレス",
     "顧客要求への対応（受注維持）。品番別不良率を導入後に確認", "5台。HTM-8709はH&H代替で▲9,480USDの余地（営業確認要）", S_SUM),
    (3, "1・2", "ロックミシン OL／IL", "必須（生産継続）",
     "OL実稼働86台（現有102−故障4−サンプル12）に対しピーク112台→不足26台。IL実稼働64台、ピーク70台→不足6台。繁忙期は第2工場から10〜12台借用、2024年10〜12月は外部レンタル約10台",
     "第2工場借用の継続：先方の受注増で困難／外部レンタル：都度D/N請求／故障4台：修理復旧不可", "Juki MO-6814・MO-6816（サーボ・自動糸切り）",
     "不足解消。ペダルOFFで即停止→曲線部の縫いズレ低減。電力1台▲469.64kWh/年", "外部レンタル回避（2024実績額を入力）＋電力37.57USD/台/年（副次）",
     "OL5台（更新4＋純増1）・IL2台", f"{S_LOCK}\n{S_SERVO}"),
    (4, "4", "パターンシーマ 中型 60×40", "改善",
     "手作業6工程・340秒、改善治具＋1本針3工程・330秒（例）。熟練者依存、習熟期間が長い。GW系はポケットファスナーをPSで縫う指示あり",
     "改善治具：工程は集約できるが作業者は全時間拘束、熟練依存が残る", "Hikari HMS-6040（既存5台、稼働問題なし）",
     "2工程・310秒（例）。自動縫製中は次製品を準備でき拘束時間が減る。全ケースでサイクルが短くなるわけではない",
     "省人工数（07で実測値から算出）、習熟期間短縮、品質安定", "3台。運用設計（08）とセットで導入", f"{S_PS}\n{S_OPS}"),
    (5, "5", "パターンシーマ 大型 130×95", "改善",
     "現有3台。2025年大型PSレンタル3,016.54USD。キルティングは外注（25FWで約3千USD、納期約1か月）。GW系はフロントファスナーをPSで縫う指示あり",
     "レンタル継続・キルティング外注継続", "Hikari HMS-13095A/RND（360度回転ヘッド）",
     "レンタル・外注を内製化、納期約1か月短縮、フロントファスナーの直線性向上", "レンタル3,016.54＋外注約3,000＝約6,000USD/年＋省人効果",
     "6→4台（段階導入、効果確認後に追加）", f"{S_PS}\n{S_RENT}"),
    (6, "6", "本縫針送りミシン", "必須（生産継続）",
     "現有17台（DUMA8・JUKI9）、各ライン1台必要（31台）。25年度の積み残し分をレンタルで対応", "Juki DDL-9010（3,500USD/台）",
     "DUMA DM1968L-NF 880USD/台", "価格は約1/4。部品はメーカー専用品（互換なし）。使用実績1年",
     "レンタル回避額は2025明細で確認（原案13,440USD/年は不整合→修正）", "10台", f"{S_RENT}\n{S_CHK}"),
    (7, "9", "温冷プレス 60×12", "必須（生産継続）",
     "30×10cm 26台で不足。面積が小さくプレス回数が増えシワのリスク", "現有機でプレス回数を増やす", "H&H CSL-201 60×12cm",
     "1回の処理面積が約2.4倍、工程秒約半分（原案）", "削減秒×対象数量で算出（数量入力待ち）", "6台", S_ROI),
    (8, "7", "シームシーリング Spec-D", "改善",
     "袖・裾など長尺パーツは限られた熟練者または2名体制で作業", "既存機へのアタッチメント：不可（H&H・工場技術で確認済）", "H&H AI-002 spec-D",
     "長尺パーツを1名で作業、テープ接着精度向上", "2名→1名分の工数削減（07で算出）", "3→2台（効果確認後に追加）", S_PS),
    (9, "10", "平プレス 125×83", "改善",
     "小型機で分割プレス→接着ムラ・シワ（TNFエアチャンバーPUR接着）", "分割プレスを継続", "H&H CS-691 125×83cm",
     "大型パーツを1回でプレス", "プレス回数削減（品番別に計測）。精度は10の方法で担保", "2→1台（検証後に追加）", S_FLAT),
    (10, "8", "シームシーリング Spec-IJ", "戦略／検証", "フード等の曲線部でシーム剥離", "－", "H&H AI-129 spec-IJ",
     "傾斜ポール切替・特殊ローラー", "剥離不良率で評価", "1台（評価用）", "－"),
    (11, "18", "超音波溶着機 US-512", "戦略／検証", "28Hz機では一部素材で溶着不良", "－", "H&H US-512（35Hz）",
     "薄手素材に対応", "サンプル試作で評価、営業提案の幅を広げる", "1台（トライアル）", "－"),
    (12, "－", "縮絨機", "戦略／検証", "生地が不安定なサプライヤーあり。過去のウール試験は効果が明確でなかった", "－", "未定", "－",
     "試験後に判断", "2026年度は計上しない", S_FELT),
    ("－", "3,16,17,\n18※,19,20", "眠り穴・自動延反機・ダウン充填機・Nawonテーピング2機種・刺繍機", "見送り",
     "－", "－", "－", "－", "－", "2026年度は計上しない（打合せ結論・原案優先②）", S_SUM),
]
header(ws, 3, ["優先", "原案No.", "設備", "投資区分", "現状方式", "改善・代替案（購入以外）", "新設備案", "差分",
               "年間効果（数値）", "投資判断（調整案）", "根拠シート"])
ws.freeze_panes = "D4"
for k, row in enumerate(cmp_rows):
    rr = 4 + k
    for i, v in enumerate(row):
        col = "ABCDEFGHIJK"[i]
        fill = PatternFill("solid", fgColor=CAT_FILL[v]) if col == "D" else None
        put(ws, f"{col}{rr}", v, fill=fill, align=CENTER if col in "ABD" else WRAP)
    ws.row_dimensions[rr].height = 88

# ================================================================ 00 表紙
ws = cover
title(ws, "APIC 2026年度設備投資 説明資料（費用対効果）",
      "対象：APIC2026年度設備投資予算案 20260403版／APIC–TI契約は2029年3月末まで延長確定")
section(ws, 3, "1. 目的と方針")
for k, t in enumerate([
    "本当に必要な設備だけに投資する。選定理由は定性的な説明だけでなく、数字で確認できる費用対効果で示す。",
    "全機種の購入を証明しようとしない。必須投資（安全・顧客要求・生産継続）／改善投資（省人・工数削減・レンタル削減・品質）／戦略・検証投資（将来案件・トライアル）に分ける。",
    "説明の流れ：現状 → 課題 → 数値根拠 → 導入効果 → 費用対効果 → 投資妥当性 → 優先順位",
    "新規条件：APICが50,000USDを負担。投資額はSuntoolを除き概ね4,000万円以下。",
]):
    put(ws, f"B{4+k}", t, border=False)
    ws.merge_cells(f"B{4+k}:E{4+k}")
    ws.row_dimensions[4 + k].height = 30
section(ws, 9, "2. 主要数値（自動計算）")
header(ws, 10, ["項目", "USD", "円", "備考"], start="B")
key = [
    ("原案（優先①）合計", f"={ref(S_SUM, f'$H${SUM_T}')}", f"={ref(S_SUM, f'$H${SUM_T}')}*{RATE}", "4,000万円を超過"),
    ("調整案 合計", f"={ref(S_SUM, f'$J${SUM_T}')}", f"={ref(S_SUM, f'$K${SUM_T}')}", "スポットプレスは暫定台数（09）"),
    ("上限（4,000万円）との差", f"={CAP}/{RATE}-C12", f"={CAP}-D12", "＋＝上限内"),
    ("APIC負担", f"={APIC}", f"={APIC}*{RATE}", "対象機種・資産保有・償却処理は要確認"),
    ("TI負担（調整案−APIC負担）", "=C12-C14", "=D12-D14", ""),
]
for k, (lab, u, y, rem) in enumerate(key):
    rr = 11 + k
    put(ws, f"B{rr}", lab, F_BOLD)
    link(ws, f"C{rr}", u, USD)
    link(ws, f"D{rr}", y, YEN)
    put(ws, f"E{rr}", rem)
for rr in (13, 15):
    ws[f"C{rr}"].font = F_BOLD
    ws[f"D{rr}"].font = F_BOLD
section(ws, 17, "3. シート構成")
header(ws, 18, ["シート", "内容", "対応する要確認事項"], start="B")
toc = [
    (S_PRM, "為替・稼働日・電気単価・償却年数・APIC負担・上限などの前提（入力）", ""),
    (S_SUM, "原案①と調整案の比較、投資区分、総額チェック、追加削減オプション", "④・新規条件"),
    (S_QA, "3/25要確認事項①〜⑦と新規条件への回答骨子・状態", "全体"),
    (S_CMP, "機種別：現状方式｜改善・代替案｜新設備案｜差分｜年間効果｜投資判断", "全体"),
    (S_LOCK, "ロックミシン必要台数（品種別標準台数→不足→借用・レンタル実績）", "①"),
    (S_SERVO, "サーボ vs クラッチの電力削減（副次効果）と品質・操作性", "①"),
    (S_PS, "パターンシーマ／Spec-Dの省人工数（Before/After実測の入力表）", "①"),
    (S_OPS, "パターンシーマの動線・人配置・縫製中の動き・測定計画", "②"),
    (S_SPOT, "スポットプレス必要台数の計算、更新分／純増分、証跡", "③"),
    (S_FLAT, "平プレスの精度検証（項目・方法・判定基準・頻度・担当）", "⑤"),
    (S_FELT, "縮絨機の検討ステップ（試験後に導入判断）", "⑦"),
    (S_RENT, "レンタル費実績と投資による回避可能額", "①"),
    (S_ROI, "改善投資の年間純効果・回収年数、当年度P/L影響（ケースA/B）", "①・新規条件"),
    (S_CHK, "原案ファイルの不整合（提出前の修正候補）", ""),
]
for k, (s, d, q) in enumerate(toc):
    rr = 19 + k
    c = put(ws, f"B{rr}", s)
    c.hyperlink = f"#'{s}'!A1"
    c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    put(ws, f"C{rr}", d)
    put(ws, f"D{rr}", q, align=CENTER)
section(ws, 34, "4. 凡例")
legend = [
    ("青字＋黄色", "入力値（工場データ・前提）。ここを変更すると全体が再計算される", F_INPUT, FILL_INPUT),
    ("黒字", "計算式", F_BASE, None),
    ("緑字", "他シートからの参照", F_LINK, None),
    ("灰色斜体", "記入例（合計に含めない）", F_EX, FILL_EX),
    ("薄緑", "主要な結果", F_BOLD, FILL_KEY),
]
for k, (lab, d, f, fill) in enumerate(legend):
    rr = 35 + k
    put(ws, f"B{rr}", lab, f, fill)
    put(ws, f"C{rr}", d)
section(ws, 41, "5. 未確定データ（工場・関係部署から入手が必要）")
pending = [
    "レンタル明細 2024・2026年（機種・型式・台数・期間・D/N・style/line）と、投資で回避できる部分",
    "パターンシーマ・Spec-DのBefore/After実測（工程秒・自動縫製秒・並行作業秒）と対象品番の年間数量",
    "スポットプレス旧型の台数・年式・修理可否、ライン別使用台数、ピーク需要、写真・ヒヤリハット",
    "平プレスの判定基準（メーカー規格）と校正周期",
    "APIC負担50,000USDの対象機種・資産保有者・償却処理",
    "1人当たり年間人件費（省人効果の金額換算用）",
    "為替レートの統一（150円／160円）",
]
for k, t in enumerate(pending):
    put(ws, f"B{42+k}", f"・{t}", border=False)
    ws.merge_cells(f"B{42+k}:E{42+k}")

for w in wb.worksheets:
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.orientation = "landscape"
    w.page_setup.paperSize = w.PAPERSIZE_A4
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0

wb.active = 0
wb.save(OUT)
print(OUT)
